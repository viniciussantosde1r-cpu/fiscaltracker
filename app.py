from datetime import date, datetime
from io import BytesIO

# Fiscal Tracker - Fase 2.1F: consolidacao da nomenclatura Report Fiscalizacoes
from urllib.parse import quote, urlencode
from email.message import EmailMessage
from email.policy import SMTP
import os

import streamlit as st

from database import (
    FASES_FISCALIZACAO,
    PRIORIDADES_EXECUTIVAS,
    STATUS_FLASH_REPORT,
    STATUS_TAREFAS,
    TIPOS_DOCUMENTO,
    adicionar_documento_fiscalizacao,
    alterar_fase_fiscalizacao,
    arquivar_fiscalizacao,
    atualizar_dados_fiscalizacao,
    atualizar_flash_report,
    atualizar_prazos_fiscalizacao,
    atualizar_status_tarefa,
    cadastrar_fiscalizacao,
    cadastrar_tarefa,
    calcular_status_prazo,
    calcular_status_prazo_tarefa,
    buscar_fiscalizacao_por_codigo,
    excluir_documento_fiscalizacao,
    inicializar_banco,
    ler_conteudo_documento,
    listar_documentos_fiscalizacao,
    listar_fiscalizacoes,
    listar_flash_report_executivo,
    listar_prazos_fiscalizacoes,
    obter_dados_flash_report,
    obter_detalhes_fiscalizacao,
    obter_resumo_prazos,
    obter_resumo_tarefas,
    reativar_fiscalizacao,
    verificar_integridade_documento,

    # Áreas & Envolvidos - Fase 2.0A2
    PAPEIS_ENVOLVIMENTO,
    listar_areas_cadastradas,
    cadastrar_area,
    atualizar_area,
    alterar_status_area,
    listar_pessoas_cadastradas,
    cadastrar_pessoa,
    atualizar_pessoa,
    alterar_status_pessoa,
    vincular_pessoa_area,
    remover_vinculo_pessoa_area,
    listar_pessoas_area,
    obter_destinatarios_area,

    # Comunicações - Fase 2.0B3
    TIPOS_COMUNICACAO,
    obter_destinatarios_fiscalizacao,
    criar_comunicacao,
    obter_comunicacao,
    listar_comunicacoes_fiscalizacao,
    atualizar_status_comunicacao,
    excluir_comunicacao_rascunho,

    # Encerramento e Aprovações - Fase 2.0D2
    listar_candidatos_aprovacao_encerramento,
    definir_aprovador_encerramento,
    remover_aprovador_encerramento,
    listar_aprovacoes_encerramento,
    registrar_ciencia_encerramento,
    revogar_ciencia_encerramento,
    obter_status_encerramento,
)
from report_generator import gerar_flash_report_pdf, gerar_dossie_final_zip


# ============================================================
# CONFIGURACAO
# ============================================================

st.set_page_config(
    page_title="Fiscal Tracker",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded",
)

try:
    inicializar_banco()
except Exception:
    st.error(
        "Nao foi possivel inicializar o banco de dados do Fiscal Tracker. "
        "Verifique o arquivo database.py e a pasta de dados."
    )
    st.stop()


TRIBUTOS_DISPONIVEIS = [
    "ICMS",
    "ISS",
    "IRPJ",
    "CSLL",
    "PIS",
    "COFINS",
    "INSS",
    "IPI",
]

AREAS_FALLBACK = [
    "Jurídico Tributário",
    "Financeiro",
    "Controladoria",
    "Auditoria Interna",
    "Diretoria",
]

PROBABILIDADES = [
    "Provável",
    "Possível",
    "Remota",
]

# ============================================================
# INTERNACIONALIZACAO - FASE 2.0D5
# ============================================================
# A camada de idioma altera apenas textos fixos de interface e exibicao.
# Valores persistidos no banco (fases, status, probabilidades e tipos)
# continuam em portugues para preservar compatibilidade e integridade.

IDIOMAS_UI = {
    "Português": "pt",
    "English": "en",
}

TRADUCOES_UI = {
    "pt": {},
    "en": {
        "SYNGENTA BRASIL · INDIRECT TAX OPERATIONS": "SYNGENTA BRAZIL · INDIRECT TAX OPERATIONS",
        "Tax Controversy Management": "Tax Controversy Management",
        "Indirect Tax Operations": "Indirect Tax Operations",
        "Navegação": "Navigation",
        "Dashboard": "Dashboard",
        "Fiscalizações": "Tax Inspections",
        "Nova fiscalização": "New Tax Inspection",
        "Áreas & Envolvidos": "Areas & Stakeholders",
        "Detalhe": "Details",
        "Comunicações": "Communications",
        "Flash Report": "Flash Report",
        "Report Fiscalizações": "Tax Inspection Report",
        "Gerar Report Fiscalizações Final": "Generate Final Tax Inspection Report",
        "Baixar Report Fiscalizações Final": "Download Final Tax Inspection Report",
        "Gerar Report Fiscalizações consolidado": "Generate Consolidated Tax Inspection Report",
        "Baixar Report Fiscalizações": "Download Tax Inspection Report",
        "Salvar Report Fiscalizações": "Save Tax Inspection Report",
        "Flash Report Executivo": "Executive Flash Report",
        "Fiscalizações de relevância executiva": "Executive-relevance tax inspections",
        "Prioridade executiva": "Executive priority",
        "Atualização executiva": "Executive update",
        "Próximo passo": "Next step",
        "Documentos": "Documents",
        "Todas as empresas": "All companies",
        "Todas as prioridades": "All priorities",
        "Todas as probabilidades": "All probabilities",
        "Todas as fases": "All phases",
        "Casos relevantes": "Relevant cases",
        "Alta prioridade": "High priority",
        "Prazos críticos": "Critical deadlines",
        "Abrir fiscalização": "Open tax inspection",
        "Nenhuma fiscalização está marcada para o Flash Report Executivo.": "No tax inspection is currently marked for the Executive Flash Report.",
        "Exportar Flash Report para Excel": "Export Flash Report to Excel",
        "Baixar Flash Report (.xlsx)": "Download Flash Report (.xlsx)",
        "Excel gerado com os filtros atuais do Dashboard.": "Excel generated with the current Dashboard filters.",
        "Portfolio Overview": "Portfolio Overview",
        "Deadlines & Critical Actions": "Deadlines & Critical Actions",
        "Internal Tasks": "Internal Tasks",
        "Upcoming Deadlines": "Upcoming Deadlines",
        "Portfolio Filters": "Portfolio Filters",
        "Tax Controversy Portfolio": "Tax Controversy Portfolio",
        "Executive Overview": "Executive Overview",
        "Case Information": "Case Information",
        "Deadlines": "Deadlines",
        "Workflow": "Workflow",
        "Stakeholders & Areas": "Stakeholders & Areas",
        "Documents & Evidence": "Documents & Evidence",
        "Timeline": "Timeline",
        "Recipients": "Recipients",
        "Prepare Communication": "Prepare Communication",
        "Outlook Draft": "Outlook Draft",
        "Communication History": "Communication History",
        "Report Management": "Report Management",
        "Governança do encerramento": "Closure Governance",
        "Comunicação final de encerramento": "Final Closure Communication",
        "Ações do documento": "Document Actions",
        "01 — Identificação": "01 — Identification",
        "02 — Executive Summary": "02 — Executive Summary",
        "03 — Inspection Background": "03 — Inspection Background",
        "04 — Tax Controversy": "04 — Tax Controversy",
        "05 — Developments": "05 — Developments",
        "06 — Deadlines": "06 — Deadlines",
        "07 — Actions Taken": "07 — Actions Taken",
        "08 — Internal Tasks": "08 — Internal Tasks",
        "09 — Documents & Evidence": "09 — Documents & Evidence",
        "10 — Risks & Impacts": "10 — Risks & Impacts",
        "11 — Next Steps": "11 — Next Steps",
        "12 — Current Position": "12 — Current Position",
        "Cadastro de Áreas": "Area Registry",
        "Áreas cadastradas": "Registered Areas",
        "Cadastro de Pessoas": "People Registry",
        "Pessoas cadastradas": "Registered People",
        "Vínculos Área × Pessoa": "Area × Person Links",
        "Mapa de Envolvidos": "Stakeholder Map",
        "Áreas": "Areas",
        "Pessoas": "People",
        "Vínculos": "Links",
        "Buscar": "Search",
        "Probabilidade": "Probability",
        "Situação": "Status",
        "Abrir": "Open",
        "← Voltar para fiscalizações": "← Back to tax inspections",
        "Exposição": "Exposure",
        "Fase atual": "Current phase",
        "Prazo": "Deadline",
        "Recebimento": "Received",
        "Prazo atual": "Current deadline",
        "Editar dados cadastrais": "Edit case information",
        "Salvar alterações": "Save changes",
        "Editar prazo": "Edit deadline",
        "Salvar prazo": "Save deadline",
        "Nova tarefa": "New task",
        "Criar tarefa": "Create task",
        "Atualizar": "Update",
        "Nova fase": "New phase",
        "Observação": "Notes",
        "Atualizar fase": "Update phase",
        "Salvar aprovador": "Save approver",
        "Adicionar ao dossiê": "Add to case file",
        "Confirmar exclusão": "Confirm deletion",
        "Excluir": "Delete",
        "Tipo da comunicação": "Communication type",
        "Destinatários *": "Recipients *",
        "Assunto *": "Subject *",
        "Mensagem *": "Message *",
        "Preparar comunicação": "Prepare communication",
        "Abrir no Outlook": "Open in Outlook",
        "Reabrir no Outlook": "Reopen in Outlook",
        "Marcar como enviada": "Mark as sent",
        "Descartar preparação": "Discard preparation",
        "Gerar Report Fiscalizações Final": "Generate Final Tax Inspection Report",
        "Baixar Report Fiscalizações Final": "Download Final Tax Inspection Report",
        "Gerar Report Fiscalizações consolidado": "Generate Consolidated Tax Inspection Report",
        "Baixar Report Fiscalizações": "Download Tax Inspection Report",
        "Gerar Dossiê Final ZIP": "Generate Final Case File ZIP",
        "Baixar Dossiê Final": "Download Final Case File",
        "Editar conteúdo do relatório": "Edit report content",
        "Salvar Report Fiscalizações": "Save Tax Inspection Report",
        "Cadastrar nova área": "Register new area",
        "Cadastrar área": "Register area",
        "Cadastrar nova pessoa": "Register new person",
        "Cadastrar pessoa": "Register person",
        "Salvar vínculo": "Save link",
        "Cadastrar fiscalização": "Register tax inspection",
        "Case Information": "Case Information",
        "Tax Information": "Tax Information",
        "Resumo executivo": "Executive summary",
        "Problema / controvérsia": "Issue / controversy",
        "Causa / fundamento": "Cause / legal basis",
        "Ações necessárias": "Required actions",
        "Ações implementadas": "Implemented actions",
        "Riscos e impactos": "Risks and impacts",
        "Próximos passos": "Next steps",
        "Conclusão / posição atual": "Conclusion / current position",
        "Não informado.": "Not provided.",
        "Nenhum prazo ativo cadastrado.": "No active deadline registered.",
        "Nenhuma fiscalização encontrada.": "No tax inspection found.",
        "Nenhuma fiscalização selecionada.": "No tax inspection selected.",
        "Fiscalização não encontrada.": "Tax inspection not found.",
        "Nenhuma tarefa cadastrada.": "No task registered.",
        "Nenhum desdobramento registrado.": "No development recorded.",
        "Nenhum documento ou evidência foi registrado no dossiê.": "No document or evidence has been registered in the case file.",
        "Fiscalização encerrada.": "Tax inspection closed.",
        "Workflow indisponível para registros arquivados.": "Workflow is unavailable for archived records.",
        "A fiscalização já está nesta fase.": "The tax inspection is already in this phase.",
        "Nenhuma alteração de status foi realizada.": "No status change was made.",
        "Dados da fiscalização atualizados com sucesso.": "Tax inspection data updated successfully.",
        "Prazo atualizado com sucesso.": "Deadline updated successfully.",
        "Tarefa criada com sucesso.": "Task created successfully.",
        "Status da tarefa atualizado.": "Task status updated.",
        "Fase atualizada com sucesso.": "Phase updated successfully.",
        "Fiscalização encerrada com governança de aprovação concluída.": "Tax inspection closed with approval governance completed.",
        "Todos os aprovadores essenciais registraram ciência. O encerramento está liberado.": "All essential approvers have acknowledged. Closure is authorized.",
        "Aprovadores cadastrados": "Registered approvers",
        "Aprovadores essenciais": "Essential approvers",
        "Essenciais pendentes": "Pending essentials",
        "Ciências essenciais": "Essential acknowledgements",
        "Pendências": "Pending items",
        "Destinatários automáticos": "Automatic recipients",
        "Comunicações enviadas": "Sent communications",
        "Pendentes de confirmação": "Pending confirmation",
        "Áreas ativas": "Active areas",
        "Pessoas ativas": "Active people",
        "ATIVA": "ACTIVE",
        "ARQUIVADA": "ARCHIVED",
        "INATIVA": "INACTIVE",
        "Ativa": "Active",
        "Inativa": "Inactive",
        "Ciente": "Acknowledged",
        "Pendente": "Pending",
        "Essencial": "Essential",
        "Aprovador": "Approver",
        "Provável": "Probable",
        "Possível": "Possible",
        "Remota": "Remote",
        "Intimação recebida": "Notice received",
        "Em análise interna": "Under internal review",
        "Resposta protocolada": "Response filed",
        "Auto de infração": "Tax assessment",
        "Recurso": "Appeal",
        "Julgamento": "Decision",
        "Encerrada": "Closed",
        "Pendente": "Pending",
        "Em andamento": "In progress",
        "Concluída": "Completed",
        "Em elaboração": "Drafting",
        "Em revisão": "Under review",
        "Finalizado": "Finalized",
        "Rascunho": "Draft",
        "Preparada": "Prepared",
        "Enviando": "Sending",
        "Enviada": "Sent",
        "Falhou": "Failed",
        "Atualização": "Update",
        "Prazo": "Deadline",
        "Solicitação de informação": "Information request",
        "Nova intimação": "New notice",
        "Resposta protocolada": "Response filed",
        "Encerramento": "Closure",
        "Outros": "Other",
        "Vencido": "Overdue",
        "Vence hoje": "Due today",
        "Crítico": "Critical",
        "Crítica": "Critical",
        "Atenção": "Attention",
        "No prazo": "On time",
        "Sem prazo": "No deadline",
    },
}


TRADUCOES_UI["pt"].update({
    "Tax Controversy Management": "Gestão de Contencioso Tributário",
    "Indirect Tax Operations": "Operações de Tributos Indiretos",
    "Portfolio Overview": "Visão Geral do Portfólio",
    "Deadlines & Critical Actions": "Prazos e Ações Críticas",
    "Internal Tasks": "Tarefas Internas",
    "Upcoming Deadlines": "Próximos Prazos",
    "Portfolio Filters": "Filtros do Portfólio",
    "Tax Controversy Portfolio": "Portfólio de Contencioso Tributário",
    "Executive Overview": "Visão Executiva",
    "Case Information": "Informações da Fiscalização",
    "Deadlines": "Prazos",
    "Workflow": "Fluxo da Fiscalização",
    "Stakeholders & Areas": "Envolvidos e Áreas",
    "Documents & Evidence": "Documentos e Evidências",
    "Timeline": "Linha do Tempo",
    "Recipients": "Destinatários",
    "Prepare Communication": "Preparar Comunicação",
    "Outlook Draft": "Rascunho no Outlook",
    "Communication History": "Histórico de Comunicações",
    "Report Management": "Gestão do Relatório",
    "02 — Executive Summary": "02 — Resumo Executivo",
    "03 — Inspection Background": "03 — Contexto da Fiscalização",
    "04 — Tax Controversy": "04 — Controvérsia Tributária",
    "05 — Developments": "05 — Desdobramentos",
    "06 — Deadlines": "06 — Prazos",
    "07 — Actions Taken": "07 — Ações Realizadas",
    "08 — Internal Tasks": "08 — Tarefas Internas",
    "09 — Documents & Evidence": "09 — Documentos e Evidências",
    "10 — Risks & Impacts": "10 — Riscos e Impactos",
    "11 — Next Steps": "11 — Próximos Passos",
    "12 — Current Position": "12 — Posição Atual",
    "Case Information": "Informações da Fiscalização",
    "Tax Information": "Informações Tributárias",
})

TRADUCOES_UI["en"].update({
    "Alta": "High",
    "Média": "Medium",
    "Baixa": "Low",
    "Fiscalizações abertas": "Open tax inspections", "Exposição total": "Total exposure",
    "Classificação provável": "Probable classification", "Prazos vencidos": "Overdue deadlines",
    "Vencidos": "Overdue", "Críticos": "Critical", "Atenção": "Attention", "Sem prazo": "No deadline",
    "Tarefas vencidas": "Overdue tasks", "Críticas": "Critical", "Pendentes": "Pending",
    "Em andamento": "In progress", "Recebimento": "Received", "Prazo atual": "Current deadline",
    "Aprovadores cadastrados": "Registered approvers", "Aprovadores essenciais": "Essential approvers",
    "Essenciais pendentes": "Pending essentials", "Ciências essenciais": "Essential acknowledgements",
    "Pendências": "Pending items", "Destinatários automáticos": "Automatic recipients",
    "Comunicações enviadas": "Sent communications", "Pendentes de confirmação": "Pending confirmation",
    "Áreas ativas": "Active areas", "Pessoas ativas": "Active people", "Vínculos": "Links",
})

def idioma_atual():
    return st.session_state.get("idioma_ui", "pt")

def t(texto):
    if texto is None:
        return texto
    return TRADUCOES_UI.get(idioma_atual(), {}).get(str(texto), str(texto))

def traduzir_valor(valor):
    return t(valor)

def traduzir_opcao(valor):
    return traduzir_valor(valor)

EXTENSOES_UPLOAD = [
    "pdf",
    "doc",
    "docx",
    "xls",
    "xlsx",
    "csv",
    "txt",
    "msg",
    "eml",
    "png",
    "jpg",
    "jpeg",
]


# ============================================================
# DESIGN SYSTEM
# ============================================================

st.markdown(
    """
    <style>
    :root {
        --syngenta-green: #006B3F;
        --syngenta-green-dark: #005533;
        --corporate-navy: #17365D;
        --deep-navy: #102A43;
        --background: #F4F6F5;
        --surface: #FFFFFF;
        --border: #C8D0CC;
        --text: #1D2924;
        --muted: #66736D;
    }

    html, body, [class*="css"] {
        font-family: Arial, Helvetica, sans-serif;
    }

    .stApp {
        background-color: var(--background);
        color: var(--text);
    }

    .block-container {
        max-width: 1500px;
        padding-top: 1.4rem;
        padding-bottom: 4rem;
    }

    section[data-testid="stSidebar"] {
        background-color: #FFFFFF;
        border-right: 1px solid var(--border);
    }

    section[data-testid="stSidebar"] h1 {
        color: var(--syngenta-green);
        font-size: 1.35rem;
        font-weight: 700;
    }

    h1 {
        color: var(--deep-navy);
        font-weight: 700;
    }

    h2, h3, h4 {
        color: var(--corporate-navy);
        font-weight: 700;
    }

    div[data-testid="stVerticalBlockBorderWrapper"] {
        background-color: var(--surface) !important;
        border: 1px solid var(--border) !important;
        border-radius: 8px !important;
        box-shadow: none !important;
    }

    div[data-testid="stMetric"] {
        background-color: #FFFFFF;
        border: 1px solid var(--border);
        border-top: 3px solid var(--syngenta-green);
        border-radius: 8px;
        padding: 14px 16px;
    }

    div[data-testid="stMetricLabel"] {
        color: var(--muted);
    }

    div[data-testid="stMetricValue"] {
        color: var(--deep-navy);
        font-weight: 700;
    }

    div.stButton > button,
    div.stDownloadButton > button {
        border: 1px solid var(--border);
        border-radius: 6px;
        min-height: 38px;
        font-weight: 700;
    }

    div.stButton > button[kind="primary"] {
        background-color: var(--syngenta-green);
        border-color: var(--syngenta-green);
        color: #FFFFFF;
    }

    div.stButton > button[kind="primary"]:hover {
        background-color: var(--syngenta-green-dark);
        border-color: var(--syngenta-green-dark);
    }

    input,
    textarea,
    div[data-baseweb="select"] > div {
        border-color: var(--border) !important;
        border-radius: 6px !important;
    }

    div[data-testid="stForm"],
    details[data-testid="stExpander"] {
        background-color: #FFFFFF;
        border: 1px solid var(--border) !important;
        border-radius: 8px !important;
    }

    div[data-testid="stForm"] {
        padding: 14px;
    }

    button[data-baseweb="tab"] {
        color: var(--muted);
        font-weight: 700;
    }

    button[data-baseweb="tab"][aria-selected="true"] {
        color: var(--corporate-navy);
    }

    div[data-baseweb="tab-highlight"] {
        background-color: var(--syngenta-green) !important;
    }

    div[data-testid="stAlert"],
    div[data-testid="stDataFrame"] {
        border: 1px solid var(--border);
        border-radius: 8px;
    }

    div[data-testid="stProgress"] > div > div {
        background-color: var(--syngenta-green);
    }

    div[data-testid="stCaptionContainer"] {
        color: var(--muted);
    }
    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# FUNCOES AUXILIARES
# ============================================================

def obter_nomes_areas_ativas():
    """
    Carrega as áreas ativas diretamente do cadastro administrativo.
    Mantém uma lista fallback apenas se houver falha inesperada de leitura.
    """
    try:
        areas_cadastradas = listar_areas_cadastradas(
            incluir_inativas=False
        )
        nomes = [
            area["nome"]
            for area in areas_cadastradas
            if area.get("ativa")
        ]

        if nomes:
            return nomes

    except Exception:
        pass

    return AREAS_FALLBACK.copy()


def formatar_moeda(valor):
    valor = valor or 0
    texto = f"{valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R$ {texto}"


def formatar_milhoes(valor):
    valor = valor or 0
    if valor >= 1_000_000:
        return f"R$ {valor / 1_000_000:.1f} mi".replace(".", ",")
    if valor >= 1_000:
        return f"R$ {valor / 1_000:.1f} mil".replace(".", ",")
    return formatar_moeda(valor)


def formatar_data(valor):
    if not valor:
        return "-"
    try:
        ano, mes, dia = str(valor)[:10].split("-")
        return f"{dia}/{mes}/{ano}"
    except Exception:
        return str(valor)


def formatar_data_hora(valor):
    if not valor:
        return "-"
    texto = str(valor)
    try:
        ano, mes, dia = texto[:10].split("-")
        if len(texto) >= 16:
            return f"{dia}/{mes}/{ano} {texto[11:16]}"
        return f"{dia}/{mes}/{ano}"
    except Exception:
        return texto


def formatar_tamanho_arquivo(tamanho_bytes):
    tamanho_bytes = tamanho_bytes or 0
    if tamanho_bytes < 1024:
        return f"{tamanho_bytes} B"
    if tamanho_bytes < 1024 * 1024:
        return f"{tamanho_bytes / 1024:.1f} KB"
    return f"{tamanho_bytes / (1024 * 1024):.1f} MB"


def texto_dias_restantes(dias):
    if dias is None:
        return "-"
    if idioma_atual() == "en":
        if dias < 0:
            atraso = abs(dias)
            return "1 day overdue" if atraso == 1 else f"{atraso} days overdue"
        if dias == 0:
            return "Due today"
        if dias == 1:
            return "1 day remaining"
        return f"{dias} days remaining"
    if dias < 0:
        atraso = abs(dias)
        return "1 dia em atraso" if atraso == 1 else f"{atraso} dias em atraso"
    if dias == 0:
        return "Vence hoje"
    if dias == 1:
        return "1 dia restante"
    return f"{dias} dias restantes"


def calcular_proxima_fase(fase_atual):
    try:
        indice = FASES_FISCALIZACAO.index(fase_atual)
    except ValueError:
        return None
    if indice >= len(FASES_FISCALIZACAO) - 1:
        return None
    return FASES_FISCALIZACAO[indice + 1]


def limpar_pdf_sessao(fiscalizacao_id):
    for chave in (
        f"flash_pdf_{fiscalizacao_id}",
        f"flash_pdf_nome_{fiscalizacao_id}",
        f"dossie_zip_{fiscalizacao_id}",
        f"dossie_zip_nome_{fiscalizacao_id}",
        f"encerramento_pdf_{fiscalizacao_id}",
        f"encerramento_pdf_nome_{fiscalizacao_id}",
    ):
        st.session_state.pop(chave, None)


def titulo_pagina(titulo, descricao):
    st.caption(t("SYNGENTA BRASIL · INDIRECT TAX OPERATIONS"))
    st.title(titulo)
    st.caption(descricao)
    st.divider()


def titulo_secao(titulo, descricao=None):
    st.subheader(titulo)
    if descricao:
        st.caption(descricao)


def exibir_texto_relatorio(texto):
    if texto and str(texto).strip():
        st.write(texto)
    else:
        st.caption(t(t(t("Não informado."))))


def exibir_status_prazo(status, dias=None):
    status_exibido = traduzir_valor(status)
    if status in ["Vencido", "Vence hoje"]:
        st.error(status_exibido)
    elif status in ["Crítico", "Crítica", "Atenção"]:
        st.warning(status_exibido)
    elif status in ["No prazo", "Concluída"]:
        st.success(status_exibido)
    else:
        st.info(status_exibido)
    if dias is not None:
        st.caption(texto_dias_restantes(dias))


def tratar_erro_operacional(erro, contexto="operação"):
    if isinstance(erro, ValueError):
        st.warning(str(erro))
    elif isinstance(erro, FileNotFoundError):
        st.error(
            "O arquivo solicitado não foi encontrado no repositório do Fiscal Tracker."
        )
    elif isinstance(erro, PermissionError):
        st.error(
            "O Fiscal Tracker não possui permissão para acessar o arquivo ou pasta necessária."
        )
    else:
        st.error(
            f"Não foi possível concluir a {contexto}. Tente novamente. "
            "Se o problema persistir, verifique o banco de dados e o repositório de documentos."
        )


def exibir_sucesso(mensagem):
    st.success(mensagem)


def validar_campos_obrigatorios(campos):
    erros = []
    for nome, valor in campos:
        if valor is None or (isinstance(valor, str) and not valor.strip()):
            erros.append(f"Informe {nome}.")
    return erros


def validar_periodo(data_inicial, data_final, nome_inicial="data inicial", nome_final="data final"):
    if data_inicial and data_final and data_final < data_inicial:
        return f"A {nome_final} não pode ser anterior à {nome_inicial}."
    return None



# ============================================================
# COMUNICACOES ASSISTIDAS
# FASE 2.0B3
# ============================================================

def montar_url_outlook_mailto(
    destinatarios,
    assunto,
    mensagem,
):
    """
    Cria um link mailto para abrir o cliente de e-mail
    padrao do computador, normalmente o Outlook em
    ambiente corporativo.

    O Fiscal Tracker nao considera o e-mail enviado
    automaticamente. O envio continua dependendo da
    confirmacao do usuario no Outlook.
    """

    emails = []

    for destinatario in destinatarios or []:

        if isinstance(
            destinatario,
            dict,
        ):

            email = (
                destinatario.get(
                    "email"
                )
                or ""
            )

        else:

            email = str(
                destinatario
            )

        email = (
            email
            .strip()
            .lower()
        )

        if (
            email
            and email not in emails
        ):

            emails.append(
                email
            )

    if not emails:

        return None

    parametros = urlencode(
        {
            "subject": (
                assunto
                or ""
            ),
            "body": (
                mensagem
                or ""
            ),
        },
        quote_via=quote,
    )

    return (
        "mailto:"
        + ",".join(
            emails
        )
        + "?"
        + parametros
    )


def gerar_rascunho_eml_com_anexo(
    destinatarios,
    assunto,
    mensagem,
    anexo_bytes,
    nome_anexo,
):
    """Gera um .eml editavel com o Report Fiscalizações Final anexado."""
    emails = []
    for destinatario in destinatarios or []:
        if isinstance(destinatario, dict):
            email = str(destinatario.get("email") or "").strip()
        else:
            email = str(destinatario or "").strip()
        if email and email.lower() not in [item.lower() for item in emails]:
            emails.append(email)

    if not emails:
        raise ValueError("A comunicação não possui destinatários válidos.")
    if not anexo_bytes:
        raise ValueError("Gere o Report Fiscalizações Final antes de preparar o e-mail.")

    msg = EmailMessage(policy=SMTP)
    msg["To"] = ", ".join(emails)
    msg["Subject"] = str(assunto or "")
    # Faz o Outlook tratar o arquivo como mensagem ainda não enviada.
    msg["X-Unsent"] = "1"
    msg.set_content(str(mensagem or ""), charset="utf-8")
    msg.add_attachment(
        bytes(anexo_bytes),
        maintype="application",
        subtype="pdf",
        filename=str(nome_anexo or "Report_Fiscalizacao_Final.pdf"),
    )
    return msg.as_bytes(policy=SMTP)


def destinatarios_para_texto(
    destinatarios
):
    """
    Formata os destinatarios para exibicao na interface.
    """

    if not destinatarios:

        return "-"

    linhas = []

    for item in destinatarios:

        nome = (
            item.get(
                "nome"
            )
            or ""
        )

        email = (
            item.get(
                "email"
            )
            or ""
        )

        papel = (
            item.get(
                "papel"
            )
            or ""
        )

        if nome:

            texto = (
                f"{nome} <{email}>"
            )

        else:

            texto = email

        if papel:

            texto += (
                f" · {papel}"
            )

        linhas.append(
            texto
        )

    return "\n".join(
        linhas
    )


def montar_mensagem_comunicacao_padrao(
    codigo, titulo, empresa, fase, prazo_resposta, tarefa_estracta, link_fiscalizacao=None,
):
    if idioma_atual() == "en":
        linhas = [
            "Dear all,", "",
            f"Please find below an update regarding tax inspection {codigo} - {titulo}.", "",
            f"Company: {empresa or '-'}",
            f"Current phase: {traduzir_valor(fase) if fase else '-'}",
            f"Current deadline: {formatar_data(prazo_resposta)}",
            f"e-Stracta task: {tarefa_estracta or '-'}", "",
            "For further details and evidence, please refer to Fiscal Tracker.",
        ]
        if link_fiscalizacao:
            linhas.extend([link_fiscalizacao, ""])
        linhas.extend(["Kind regards,", "Tax Controversy"])
        return "\n".join(linhas)

    linhas = [
        "Prezados,", "",
        f"Segue atualização referente à fiscalização {codigo} - {titulo}.", "",
        f"Empresa: {empresa or '-'}",
        f"Fase atual: {fase or '-'}",
        f"Prazo atual: {formatar_data(prazo_resposta)}",
        f"Tarefa e-Stracta: {tarefa_estracta or '-'}", "",
        "Para detalhes e evidências, consulte o Fiscal Tracker.",
    ]
    if link_fiscalizacao:
        linhas.extend([link_fiscalizacao, ""])
    linhas.extend(["Atenciosamente,", "Tax Controversy"])
    return "\n".join(linhas)


def montar_mensagem_encerramento_padrao(codigo, titulo, empresa, link_fiscalizacao=None):
    if idioma_atual() == "en":
        linhas = [
            "Dear all,", "",
            f"We inform you that tax inspection {codigo} - {titulo} has been closed.",
            "", f"Company: {empresa or '-'}", "Status: Closed", "",
            "The closure was completed after the mandatory acknowledgement workflow recorded in Fiscal Tracker.",
            "The Final Tax Inspection Report consolidates the main developments, approvals, documents and evidence for the case.",
            "", "The Final Tax Inspection Report is attached to the Outlook draft before sending.",
        ]
        if link_fiscalizacao:
            linhas.extend(["", "Access the case file in Fiscal Tracker:", link_fiscalizacao])
        linhas.extend(["", "Kind regards,", "Tax Controversy"])
        return "\n".join(linhas)

    linhas = [
        "Prezados,", "", f"Informamos o encerramento da fiscalização {codigo} - {titulo}.",
        "", f"Empresa: {empresa or '-'}", "Status: Encerrada", "",
        "O encerramento foi concluído após o fluxo de ciências obrigatórias registrado no Fiscal Tracker.",
        "O Report Fiscalizações Final consolida os principais desdobramentos, aprovações, documentos e evidências do caso.",
        "", "O Report Fiscalizações Final segue anexado à comunicação antes do envio.",
    ]
    if link_fiscalizacao:
        linhas.extend(["", "Acesso ao dossiê no Fiscal Tracker:", link_fiscalizacao])
    linhas.extend(["", "Atenciosamente,", "Tax Controversy"])
    return "\n".join(linhas)


# ============================================================
# ACESSO DA EQUIPE / LINKS DIRETOS
# FASE 2.0C1
# ============================================================

BASE_URL_PADRAO = "http://localhost:8501"


def obter_url_base_aplicacao():
    """
    Retorna o endereco base usado nos links compartilhaveis.

    Em producao, configure a variavel de ambiente:

        FISCAL_TRACKER_BASE_URL

    Exemplo:

        https://fiscaltracker.interno

    Enquanto nao houver publicacao interna, o sistema usa
    http://localhost:8501.
    """

    url = (
        os.getenv(
            "FISCAL_TRACKER_BASE_URL",
            "",
        )
        .strip()
        .rstrip("/")
    )

    if not url:
        url = BASE_URL_PADRAO

    return url


def montar_link_fiscalizacao(codigo, secao=None):
    """Monta o link direto de uma fiscalizacao e, opcionalmente, de uma secao."""

    codigo = str(codigo or "").strip()
    secao = str(secao or "").strip().lower()

    if not codigo:
        return obter_url_base_aplicacao()

    url = f"{obter_url_base_aplicacao()}/?case={quote(codigo)}"
    if secao:
        url += f"&section={quote(secao)}"
    return url


def gerar_excel_flash_report_executivo(itens, filtros=None):
    """Gera o Flash Report Executivo em XLSX com os filtros ativos do Dashboard."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:
        raise RuntimeError(
            "A exportação para Excel requer o pacote openpyxl. "
            "Instale com: pip install openpyxl"
        ) from exc

    itens = list(itens or [])
    filtros = dict(filtros or {})
    idioma = idioma_atual()

    wb = Workbook()
    ws = wb.active
    ws.title = "Flash Report"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A7"

    verde = "006B3F"
    verde_claro = "E8F3EE"
    navy = "17365D"
    deep_navy = "102A43"
    cinza = "667085"
    cinza_claro = "F2F4F7"
    branco = "FFFFFF"
    laranja = "F79009"
    vermelho_claro = "FEE4E2"
    verde_status = "ECFDF3"

    ws.merge_cells("A1:P1")
    c = ws["A1"]
    c.value = (
        "FISCAL TRACKER — FLASH REPORT EXECUTIVO"
        if idioma == "pt"
        else "FISCAL TRACKER — EXECUTIVE FLASH REPORT"
    )
    c.font = Font(name="Aptos Display", size=18, bold=True, color=branco)
    c.fill = PatternFill("solid", fgColor=navy)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:P2")
    c = ws["A2"]
    c.value = (
        (
            "Fiscalizações de relevância executiva | Extração: "
            if idioma == "pt"
            else "Executive-relevance tax inspections | Extracted: "
        )
        + datetime.now().strftime('%d/%m/%Y %H:%M')
    )
    c.font = Font(name="Aptos", size=10, color=cinza)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 20

    filtros_texto = []
    nomes_filtros = [
        (("Busca" if idioma == "pt" else "Search"), filtros.get("busca")),
        (("Empresa" if idioma == "pt" else "Company"), filtros.get("empresa")),
        (("Prioridade" if idioma == "pt" else "Priority"), filtros.get("prioridade")),
        (("Probabilidade" if idioma == "pt" else "Probability"), filtros.get("probabilidade")),
        (("Fase" if idioma == "pt" else "Phase"), filtros.get("fase")),
    ]
    for nome, valor in nomes_filtros:
        valor = str(valor or "").strip()
        if valor:
            filtros_texto.append(f"{nome}: {valor}")
    if not filtros_texto:
        filtros_texto.append("Filtros: Todos" if idioma == "pt" else "Filters: All")

    ws.merge_cells("A3:P3")
    c = ws["A3"]
    c.value = " | ".join(filtros_texto)
    c.font = Font(name="Aptos", size=10, italic=True, color=deep_navy)
    c.fill = PatternFill("solid", fgColor=verde_claro)
    c.alignment = Alignment(vertical="center")

    total = len(itens)
    exposicao = sum(float(item.get("exposicao") or 0) for item in itens)
    alta = sum(1 for item in itens if item.get("prioridade_executiva") == "Alta")
    criticos = sum(
        1 for item in itens
        if item.get("status_prazo") in {"Vencido", "Vence hoje", "Crítico"}
    )
    kpis = [
        (("Casos relevantes" if idioma == "pt" else "Relevant cases"), total),
        (("Exposição total (R$)" if idioma == "pt" else "Total exposure (BRL)"), exposicao),
        (("Alta prioridade" if idioma == "pt" else "High priority"), alta),
        (("Prazos críticos" if idioma == "pt" else "Critical deadlines"), criticos),
    ]
    for idx, (rotulo, valor) in enumerate(kpis):
        col_ini = 1 + idx * 4
        col_fim = col_ini + 3
        ws.merge_cells(start_row=4, start_column=col_ini, end_row=4, end_column=col_fim)
        k = ws.cell(4, col_ini)
        k.value = f"{rotulo}: {valor}"
        if rotulo in {"Exposição total (R$)", "Total exposure (BRL)"}:
            k.value = valor
            k.number_format = 'R$ #,##0.00'
        k.font = Font(name="Aptos", size=10, bold=True, color=deep_navy)
        k.fill = PatternFill("solid", fgColor=cinza_claro)
        k.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[4].height = 24

    headers_pt = [
        "Código", "Fiscalização", "Empresa", "Órgão", "Tributos",
        "Prioridade executiva", "Fase atual", "Exposição (R$)",
        "Probabilidade", "Prazo", "Situação", "Atualização executiva",
        "Próximo passo", "Documentos", "Abrir fiscalização",
        "Documentos & Evidências",
    ]
    headers_en = [
        "Code", "Tax inspection", "Company", "Tax authority", "Taxes",
        "Executive priority", "Current phase", "Exposure (BRL)",
        "Probability", "Deadline", "Status", "Executive update",
        "Next step", "Documents", "Open tax inspection",
        "Documents & Evidence",
    ]
    headers = headers_pt if idioma == "pt" else headers_en
    header_row = 6
    for col, header in enumerate(headers, 1):
        cell = ws.cell(header_row, col, header)
        cell.font = Font(name="Aptos", size=10, bold=True, color=branco)
        cell.fill = PatternFill("solid", fgColor=verde)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 32

    for row_idx, item in enumerate(itens, header_row + 1):
        valores = [
            item.get("codigo") or "-",
            item.get("titulo") or "-",
            item.get("empresa") or "-",
            item.get("orgao") or "-",
            item.get("tributos") or "-",
            traduzir_valor(item.get("prioridade_executiva")) if item.get("prioridade_executiva") else "-",
            traduzir_valor(item.get("fase")) if item.get("fase") else "-",
            float(item.get("exposicao") or 0),
            traduzir_valor(item.get("probabilidade")) if item.get("probabilidade") else "-",
            (
                date.fromisoformat(item.get("prazo_resposta"))
                if item.get("prazo_resposta")
                else None
            ),
            traduzir_valor(item.get("status_prazo")) if item.get("status_prazo") else "-",
            (item.get("texto_executivo") or "").strip() or "-",
            (item.get("proximos_passos") or "").strip() or "-",
            int(item.get("quantidade_documentos") or 0),
            ("Abrir fiscalização" if idioma == "pt" else "Open tax inspection"),
            ("Abrir documentos" if idioma == "pt" else "Open documents"),
        ]
        for col_idx, valor in enumerate(valores, 1):
            cell = ws.cell(row_idx, col_idx, valor)
            cell.font = Font(name="Aptos", size=10, color="000000")
            cell.alignment = Alignment(vertical="top", wrap_text=col_idx in {2, 4, 5, 12, 13})

        ws.cell(row_idx, 8).number_format = 'R$ #,##0.00'
        ws.cell(row_idx, 10).number_format = 'dd/mm/yyyy'

        codigo = item.get("codigo") or ""
        if codigo:
            link_caso = montar_link_fiscalizacao(codigo)
            link_docs = montar_link_fiscalizacao(codigo, "documentos")
            ws.cell(row_idx, 15).hyperlink = link_caso
            ws.cell(row_idx, 16).hyperlink = link_docs
            for col_idx in (15, 16):
                cell = ws.cell(row_idx, col_idx)
                cell.font = Font(name="Aptos", size=10, color="0563C1", underline="single")
                cell.alignment = Alignment(horizontal="center", vertical="top")

        prioridade = item.get("prioridade_executiva")
        if prioridade == "Alta":
            ws.cell(row_idx, 6).fill = PatternFill("solid", fgColor=vermelho_claro)
            ws.cell(row_idx, 6).font = Font(name="Aptos", size=10, bold=True, color="B42318")
        elif prioridade == "Média":
            ws.cell(row_idx, 6).fill = PatternFill("solid", fgColor="FFF4E5")
            ws.cell(row_idx, 6).font = Font(name="Aptos", size=10, bold=True, color=laranja)
        else:
            ws.cell(row_idx, 6).fill = PatternFill("solid", fgColor=verde_status)

        status = item.get("status_prazo")
        if status in {"Vencido", "Vence hoje", "Crítico"}:
            ws.cell(row_idx, 11).fill = PatternFill("solid", fgColor=vermelho_claro)
            ws.cell(row_idx, 11).font = Font(name="Aptos", size=10, bold=True, color="B42318")

    last_row = max(header_row + 1, header_row + len(itens))
    if itens:
        ref = f"A{header_row}:P{last_row}"
        tabela = Table(displayName="TabelaFlashReportExecutivo", ref=ref)
        tabela.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium4",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(tabela)

    larguras = {
        "A": 13, "B": 34, "C": 24, "D": 24, "E": 22, "F": 19,
        "G": 22, "H": 18, "I": 17, "J": 14, "K": 16, "L": 48,
        "M": 42, "N": 13, "O": 20, "P": 23,
    }
    for col, width in larguras.items():
        ws.column_dimensions[col].width = width

    thin = Side(style="thin", color="D0D5DD")
    for row in ws.iter_rows(min_row=header_row, max_row=last_row, min_col=1, max_col=16):
        for cell in row:
            cell.border = Border(bottom=thin)

    ws.auto_filter.ref = f"A{header_row}:P{last_row}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = f"1:{header_row}"
    ws.oddFooter.center.text = "Fiscal Tracker | Tax Controversy"
    ws.oddFooter.right.text = ("Página &P de &N" if idioma == "pt" else "Page &P of &N")

    arquivo = BytesIO()
    wb.save(arquivo)
    arquivo.seek(0)
    return arquivo.getvalue()


def limpar_case_query_param():
    """Remove os parametros de navegacao direta da fiscalizacao."""

    try:
        if "case" in st.query_params:
            del st.query_params["case"]
        if "section" in st.query_params:
            del st.query_params["section"]
    except Exception:
        pass


def definir_case_query_param(codigo, secao=None):
    """Mantem a URL sincronizada com a fiscalizacao e a secao desejada."""

    try:
        st.query_params["case"] = codigo
        secao = str(secao or "").strip().lower()
        if secao:
            st.query_params["section"] = secao
        elif "section" in st.query_params:
            del st.query_params["section"]
    except Exception:
        pass


def processar_link_direto():
    """
    Interpreta ?case=FT-XXXX e abre diretamente a fiscalizacao.

    O parametro e processado apenas quando muda, evitando
    redirecionamentos repetidos em cada rerun do Streamlit.
    """

    try:
        codigo_parametro = st.query_params.get("case")
        secao_parametro = st.query_params.get("section")
    except Exception:
        codigo_parametro = None
        secao_parametro = None

    if isinstance(
        codigo_parametro,
        list,
    ):
        codigo_parametro = (
            codigo_parametro[0]
            if codigo_parametro
            else None
        )

    codigo_parametro = str(codigo_parametro or "").strip()
    if isinstance(secao_parametro, list):
        secao_parametro = secao_parametro[0] if secao_parametro else None
    secao_parametro = str(secao_parametro or "").strip().lower()
    if secao_parametro not in {"", "documentos"}:
        secao_parametro = ""

    token_parametro = f"{codigo_parametro}|{secao_parametro}"
    codigo_processado = st.session_state.get("case_param_processado") or ""

    if not codigo_parametro:
        st.session_state[
            "case_param_processado"
        ] = ""
        return

    if token_parametro == codigo_processado:
        return

    st.session_state["case_param_processado"] = token_parametro
    st.session_state["foco_detalhe"] = secao_parametro

    try:
        fiscalizacao = buscar_fiscalizacao_por_codigo(
            codigo_parametro
        )
    except Exception:
        fiscalizacao = None

    if fiscalizacao:
        st.session_state[
            "fiscalizacao_selecionada"
        ] = fiscalizacao[0]

        st.session_state[
            "pagina"
        ] = "Detalhe"

        st.session_state[
            "menu_principal"
        ] = "Fiscalizações"
    else:
        st.session_state[
            "link_direto_invalido"
        ] = codigo_parametro


# ============================================================
# ESTADO E NAVEGACAO
# ============================================================

st.session_state.setdefault("idioma_ui", "pt")
st.session_state.setdefault("pagina", "Dashboard")
st.session_state.setdefault("menu_principal", "Dashboard")
st.session_state.setdefault("fiscalizacao_selecionada", None)
st.session_state.setdefault("case_param_processado", "")
st.session_state.setdefault("link_direto_invalido", None)
st.session_state.setdefault("foco_detalhe", "")

processar_link_direto()


def navegar_pelo_menu():
    st.session_state["pagina"] = st.session_state["menu_principal"]
    st.session_state["fiscalizacao_selecionada"] = None
    st.session_state["case_param_processado"] = ""
    st.session_state["link_direto_invalido"] = None
    st.session_state["foco_detalhe"] = ""
    limpar_case_query_param()


def voltar_para_fiscalizacoes():
    """
    Callback de navegação para retornar à lista de fiscalizações.

    O callback é executado antes da nova renderização do Streamlit,
    permitindo atualizar com segurança o estado do widget
    `menu_principal`.
    """

    st.session_state["pagina"] = "Fiscalizações"
    st.session_state["menu_principal"] = "Fiscalizações"
    st.session_state["fiscalizacao_selecionada"] = None
    st.session_state["case_param_processado"] = ""
    st.session_state["link_direto_invalido"] = None
    st.session_state["foco_detalhe"] = ""
    limpar_case_query_param()


with st.sidebar:
    st.title("Fiscal Tracker")
    st.caption(t("Tax Controversy Management"))
    idioma_escolhido = st.selectbox(
        "Idioma / Language",
        options=list(IDIOMAS_UI.keys()),
        index=0 if idioma_atual() == "pt" else 1,
        key="seletor_idioma_ui",
    )
    novo_idioma = IDIOMAS_UI[idioma_escolhido]
    if novo_idioma != idioma_atual():
        st.session_state["idioma_ui"] = novo_idioma
        st.rerun()
    st.divider()
    st.radio(
        t("Navegação"),
        [
            "Dashboard",
            "Fiscalizações",
            "Nova fiscalização",
            "Áreas & Envolvidos",
        ],
        key="menu_principal",
        on_change=navegar_pelo_menu,
        format_func=lambda opcao: t(opcao),
        label_visibility="collapsed",
    )
    st.divider()
    st.caption("SYNGENTA BRASIL")
    st.caption(t("Indirect Tax Operations"))


# ============================================================
# DASHBOARD
# ============================================================

if st.session_state.get("link_direto_invalido"):
    st.warning(
        (
            "O link direto informado não corresponde a uma "
            "fiscalização cadastrada: "
            f"{st.session_state['link_direto_invalido']}"
        )
    )
    st.session_state["link_direto_invalido"] = None

if st.session_state.pagina == "Dashboard":
    titulo_pagina(
        "Fiscal Tracker",
        ("Gestão executiva do portfólio de fiscalizações e Tax Controversy." if idioma_atual() == "pt" else "Executive management of the tax inspection and Tax Controversy portfolio."),
    )

    try:
        fiscalizacoes = listar_fiscalizacoes("Ativas")
        flash_executivo = listar_flash_report_executivo()
        resumo_prazos = obter_resumo_prazos()
        resumo_tarefas = obter_resumo_tarefas()
        prazos = listar_prazos_fiscalizacoes()
    except Exception as erro:
        tratar_erro_operacional(erro, "consulta do Dashboard")
        fiscalizacoes = []
        flash_executivo = []
        prazos = []
        resumo_prazos = {
            "vencidos": 0,
            "vence_hoje": 0,
            "criticos": 0,
            "atencao": 0,
            "no_prazo": 0,
            "sem_prazo": 0,
            "encerradas": 0,
        }
        resumo_tarefas = {
            "pendentes": 0,
            "em_andamento": 0,
            "concluidas": 0,
            "vencidas": 0,
            "vence_hoje": 0,
            "criticas": 0,
        }

    total = len(fiscalizacoes)
    encerradas = sum(1 for item in fiscalizacoes if item[8] == "Encerrada")
    abertas = total - encerradas
    exposicao_total = sum(item[6] or 0 for item in fiscalizacoes)
    provaveis = sum(1 for item in fiscalizacoes if item[7] == "Provável")

    titulo_secao(t("Portfolio Overview"))
    c1, c2, c3, c4 = st.columns(4)
    c1.metric(t("Fiscalizações abertas"), abertas)
    c2.metric(t("Exposição total"), formatar_milhoes(exposicao_total))
    c3.metric(t("Classificação provável"), provaveis)
    c4.metric(t("Prazos vencidos"), resumo_prazos["vencidos"])

    # ========================================================
    # FLASH REPORT EXECUTIVO - FASE 2.1C
    # ========================================================
    st.write("")
    titulo_secao(
        t("Flash Report Executivo"),
        (
            "Resumo das fiscalizações marcadas como relevantes para reporting executivo."
            if idioma_atual() == "pt"
            else "Summary of tax inspections marked as relevant for executive reporting."
        ),
    )

    if not flash_executivo:
        st.info(t("Nenhuma fiscalização está marcada para o Flash Report Executivo."))
    else:
        empresas_flash = sorted({
            item["empresa"] for item in flash_executivo if item.get("empresa")
        })
        fases_flash = [
            fase for fase in FASES_FISCALIZACAO
            if any(item.get("fase") == fase for item in flash_executivo)
        ]

        ff1, ff2, ff3, ff4, ff5 = st.columns([2.2, 1.5, 1.2, 1.2, 1.4])
        with ff1:
            busca_flash = st.text_input(
                t("Buscar"),
                key="flash_exec_busca",
                placeholder=(
                    "Código, fiscalização, empresa ou órgão..."
                    if idioma_atual() == "pt"
                    else "Code, tax inspection, company or tax authority..."
                ),
            )
        with ff2:
            empresa_flash = st.selectbox(
                t("Empresa"),
                [""] + empresas_flash,
                key="flash_exec_empresa",
                format_func=lambda valor: valor or t("Todas as empresas"),
            )
        with ff3:
            prioridade_flash = st.selectbox(
                t("Prioridade executiva"),
                [""] + PRIORIDADES_EXECUTIVAS,
                key="flash_exec_prioridade",
                format_func=lambda valor: traduzir_valor(valor) if valor else t("Todas as prioridades"),
            )
        with ff4:
            probabilidade_flash = st.selectbox(
                t("Probabilidade"),
                [""] + PROBABILIDADES,
                key="flash_exec_probabilidade",
                format_func=lambda valor: traduzir_valor(valor) if valor else t("Todas as probabilidades"),
            )
        with ff5:
            fase_flash = st.selectbox(
                t("Fase atual"),
                [""] + fases_flash,
                key="flash_exec_fase",
                format_func=lambda valor: traduzir_valor(valor) if valor else t("Todas as fases"),
            )

        flash_filtrado = []
        termo_flash = (busca_flash or "").strip().lower()
        for item in flash_executivo:
            texto_busca_flash = " ".join([
                str(item.get("codigo") or ""),
                str(item.get("titulo") or ""),
                str(item.get("empresa") or ""),
                str(item.get("orgao") or ""),
                str(item.get("tributos") or ""),
                str(item.get("texto_executivo") or ""),
            ]).lower()
            if termo_flash and termo_flash not in texto_busca_flash:
                continue
            if empresa_flash and item.get("empresa") != empresa_flash:
                continue
            if prioridade_flash and item.get("prioridade_executiva") != prioridade_flash:
                continue
            if probabilidade_flash and item.get("probabilidade") != probabilidade_flash:
                continue
            if fase_flash and item.get("fase") != fase_flash:
                continue
            flash_filtrado.append(item)

        total_flash = len(flash_filtrado)
        exposicao_flash = sum(item.get("exposicao") or 0 for item in flash_filtrado)
        alta_flash = sum(
            1 for item in flash_filtrado
            if item.get("prioridade_executiva") == "Alta"
        )
        criticos_flash = sum(
            1 for item in flash_filtrado
            if item.get("status_prazo") in {"Vencido", "Vence hoje", "Crítico"}
        )

        fm1, fm2, fm3, fm4 = st.columns(4)
        fm1.metric(t("Casos relevantes"), total_flash)
        fm2.metric(t("Exposição total"), formatar_milhoes(exposicao_flash))
        fm3.metric(t("Alta prioridade"), alta_flash)
        fm4.metric(t("Prazos críticos"), criticos_flash)

        if not flash_filtrado:
            st.info(
                "Nenhum caso corresponde aos filtros selecionados."
                if idioma_atual() == "pt"
                else "No case matches the selected filters."
            )
        else:
            linhas_flash = []
            for item in flash_filtrado:
                texto_executivo = (item.get("texto_executivo") or "").strip()
                proximos_passos = (item.get("proximos_passos") or "").strip()
                linhas_flash.append({
                    t("Código"): item.get("codigo") or "-",
                    t("Fiscalização"): item.get("titulo") or "-",
                    t("Empresa"): item.get("empresa") or "-",
                    ("Órgão" if idioma_atual() == "pt" else "Tax authority"): item.get("orgao") or "-",
                    ("Tributos" if idioma_atual() == "pt" else "Taxes"): item.get("tributos") or "-",
                    t("Prioridade executiva"): traduzir_valor(item.get("prioridade_executiva")),
                    t("Fase atual"): traduzir_valor(item.get("fase")),
                    t("Exposição"): formatar_moeda(item.get("exposicao")),
                    t("Probabilidade"): traduzir_valor(item.get("probabilidade")),
                    t("Prazo"): formatar_data(item.get("prazo_resposta")),
                    t("Situação"): traduzir_valor(item.get("status_prazo")),
                    t("Atualização executiva"): texto_executivo or "-",
                    t("Próximo passo"): proximos_passos or "-",
                    t("Documentos"): item.get("quantidade_documentos", 0),
                })

            st.dataframe(
                linhas_flash,
                use_container_width=True,
                hide_index=True,
                height=min(520, 38 * len(linhas_flash) + 45),
            )

            st.caption(
                (
                    "A tabela usa a Atualização Executiva; quando o campo está vazio, "
                    "o Resumo Executivo do Report Fiscalizações é utilizado automaticamente."
                )
                if idioma_atual() == "pt"
                else (
                    "The table uses the Executive Update; when it is empty, the Executive "
                    "Summary from the Tax Inspection Report is used automatically."
                )
            )

            filtros_excel_flash = {
                "busca": busca_flash,
                "empresa": empresa_flash,
                "prioridade": prioridade_flash,
                "probabilidade": probabilidade_flash,
                "fase": fase_flash,
            }
            try:
                excel_flash_bytes = gerar_excel_flash_report_executivo(
                    flash_filtrado,
                    filtros=filtros_excel_flash,
                )
                nome_excel_flash = (
                    "Flash_Report_Executivo_"
                    f"{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
                )
                ex1, ex2 = st.columns([2.2, 5.8])
                with ex1:
                    st.download_button(
                        t("Baixar Flash Report (.xlsx)"),
                        data=excel_flash_bytes,
                        file_name=nome_excel_flash,
                        mime=(
                            "application/vnd.openxmlformats-officedocument."
                            "spreadsheetml.sheet"
                        ),
                        key="flash_exec_download_excel",
                        use_container_width=True,
                    )
                with ex2:
                    st.caption(t("Excel gerado com os filtros atuais do Dashboard."))
            except Exception as erro_excel:
                st.warning(
                    (
                        f"Não foi possível gerar o Excel do Flash Report: {erro_excel}"
                        if idioma_atual() == "pt"
                        else f"The Flash Report Excel could not be generated: {erro_excel}"
                    )
                )

            opcoes_abrir_flash = {
                f"{item['codigo']} — {item['titulo']}": item
                for item in flash_filtrado
            }
            nav1, nav2, nav3 = st.columns([4.8, 1.55, 1.9])
            with nav1:
                selecao_flash = st.selectbox(
                    t("Fiscalização"),
                    list(opcoes_abrir_flash.keys()),
                    key="flash_exec_abrir_caso",
                    label_visibility="collapsed",
                )
            with nav2:
                if st.button(
                    t("Abrir fiscalização"),
                    key="flash_exec_btn_abrir",
                    type="primary",
                    use_container_width=True,
                ):
                    caso_flash = opcoes_abrir_flash[selecao_flash]
                    st.session_state["fiscalizacao_selecionada"] = caso_flash["id"]
                    st.session_state["pagina"] = "Detalhe"
                    st.session_state["foco_detalhe"] = ""
                    st.session_state["case_param_processado"] = f"{caso_flash['codigo']}|"
                    definir_case_query_param(caso_flash["codigo"])
                    st.rerun()
            with nav3:
                rotulo_documentos = (
                    "Documentos & Evidências"
                    if idioma_atual() == "pt"
                    else "Documents & Evidence"
                )
                if st.button(
                    rotulo_documentos,
                    key="flash_exec_btn_documentos",
                    use_container_width=True,
                ):
                    caso_flash = opcoes_abrir_flash[selecao_flash]
                    st.session_state["fiscalizacao_selecionada"] = caso_flash["id"]
                    st.session_state["pagina"] = "Detalhe"
                    st.session_state["foco_detalhe"] = "documentos"
                    st.session_state["case_param_processado"] = f"{caso_flash['codigo']}|documentos"
                    definir_case_query_param(caso_flash["codigo"], "documentos")
                    st.rerun()

    st.write("")
    titulo_secao(
        t("Deadlines & Critical Actions"),
        "Visão dos principais riscos operacionais relacionados a prazo.",
    )
    p1, p2, p3, p4 = st.columns(4)
    p1.metric(t("Vencidos"), resumo_prazos["vencidos"])
    p2.metric(t("Críticos"), resumo_prazos["criticos"])
    p3.metric(t("Atenção"), resumo_prazos["atencao"])
    p4.metric(t("Sem prazo"), resumo_prazos["sem_prazo"])

    if resumo_prazos["vence_hoje"] > 0:
        st.error(f"{resumo_prazos['vence_hoje']} fiscalização(ões) vencem hoje.")

    st.write("")
    titulo_secao(t("Internal Tasks"))
    t1, t2, t3, t4 = st.columns(4)
    t1.metric(t("Tarefas vencidas"), resumo_tarefas["vencidas"])
    t2.metric(t("Críticas"), resumo_tarefas["criticas"])
    t3.metric(t("Pendentes"), resumo_tarefas["pendentes"])
    t4.metric(t("Em andamento"), resumo_tarefas["em_andamento"])

    st.write("")
    titulo_secao(t("Upcoming Deadlines"), "Próximos prazos das fiscalizações ativas." if idioma_atual() == "pt" else "Upcoming deadlines for active tax inspections.")

    prazos_relevantes = [
        item
        for item in prazos
        if item["status_prazo"] not in ["Sem prazo", "Encerrada"]
    ]

    if not prazos_relevantes:
        st.info(t(t("Nenhum prazo ativo cadastrado.")))
    else:
        for item in prazos_relevantes[:10]:
            with st.container(border=True):
                pc1, pc2, pc3, pc4 = st.columns([1, 4, 2, 2])
                with pc1:
                    st.caption("CÓDIGO")
                    st.write(item["codigo"])
                with pc2:
                    st.caption("FISCALIZAÇÃO")
                    st.write(item["titulo"])
                    st.caption(item["empresa"])
                with pc3:
                    st.caption("PRAZO")
                    st.write(formatar_data(item["prazo_resposta"]))
                    st.caption(texto_dias_restantes(item["dias_restantes"]))
                with pc4:
                    st.caption("SITUAÇÃO")
                    exibir_status_prazo(item["status_prazo"])


# ============================================================
# FISCALIZACOES
# ============================================================

elif st.session_state.pagina == "Fiscalizações":
    titulo_pagina(
        t("Fiscalizações"),
        ("Consulta e gestão do portfólio de Tax Controversy." if idioma_atual() == "pt" else "Review and manage the Tax Controversy portfolio."),
    )

    titulo_secao(t("Portfolio Filters"))
    f1, f2, f3 = st.columns([2, 1, 1])
    with f1:
        busca = st.text_input(
            t("Buscar"),
            placeholder=("Código, título, empresa ou órgão..." if idioma_atual() == "pt" else "Code, title, company or tax authority..."),
        )
    with f2:
        filtro_probabilidade = st.selectbox(
            t("Probabilidade"),
            ["Todas", "Provável", "Possível", "Remota"],
        )
    with f3:
        filtro_situacao = st.selectbox(
            "Situação",
            ["Ativas", "Arquivadas", "Todas"],
        )

    try:
        fiscalizacoes = listar_fiscalizacoes(filtro_situacao)
        prazos = listar_prazos_fiscalizacoes(incluir_arquivadas=True)
    except Exception as erro:
        tratar_erro_operacional(erro, "consulta do portfólio")
        fiscalizacoes = []
        prazos = []

    mapa_prazos = {item["id"]: item for item in prazos}
    fiscalizacoes_filtradas = []

    for fiscalizacao in fiscalizacoes:
        id_, codigo, titulo, empresa, filial, orgao, exposicao, probabilidade, fase, data_entrada = fiscalizacao
        texto_busca = f"{codigo} {titulo} {empresa} {orgao}".lower()
        atende_busca = not busca or busca.lower() in texto_busca
        atende_probabilidade = filtro_probabilidade == "Todas" or probabilidade == filtro_probabilidade
        if atende_busca and atende_probabilidade:
            fiscalizacoes_filtradas.append(fiscalizacao)

    st.write("")
    titulo_secao(
        t("Tax Controversy Portfolio"),
        f"{len(fiscalizacoes_filtradas)} registro(s) encontrado(s).",
    )

    if not fiscalizacoes_filtradas:
        st.info(t(t("Nenhuma fiscalização encontrada.")))
    else:
        for fiscalizacao in fiscalizacoes_filtradas:
            id_, codigo, titulo, empresa, filial, orgao, exposicao, probabilidade, fase, data_entrada = fiscalizacao

            try:
                detalhes_card = obter_detalhes_fiscalizacao(id_)
                esta_arquivada = detalhes_card["arquivamento"]["arquivada"]
            except Exception:
                esta_arquivada = False

            prazo = mapa_prazos.get(id_)

            with st.container(border=True):
                cc1, cc2, cc3, cc4 = st.columns([4, 2, 2, 1])
                with cc1:
                    st.caption(codigo)
                    st.markdown(f"### {titulo}")
                    st.write(empresa)
                    detalhes_texto = [item for item in [filial, orgao] if item]
                    if detalhes_texto:
                        st.caption(" · ".join(detalhes_texto))
                    st.caption((t("ARQUIVADA") if esta_arquivada else t("ATIVA")) + f" · {traduzir_valor(fase)}")
                with cc2:
                    st.caption("EXPOSIÇÃO")
                    st.markdown(f"### {formatar_milhoes(exposicao)}")
                    st.caption(f"{t('Probabilidade')}: {traduzir_valor(probabilidade)}")
                with cc3:
                    st.caption("PRAZO")
                    if prazo:
                        st.write(formatar_data(prazo["prazo_resposta"]))
                        st.caption(texto_dias_restantes(prazo["dias_restantes"]))
                        exibir_status_prazo(prazo["status_prazo"])
                with cc4:
                    st.write("")
                    if st.button(
                        t("Abrir"),
                        key=f"abrir_{id_}",
                        type="primary",
                        use_container_width=True,
                    ):
                        st.session_state["fiscalizacao_selecionada"] = id_
                        st.session_state["pagina"] = "Detalhe"
                        st.session_state["case_param_processado"] = codigo
                        definir_case_query_param(codigo)
                        st.rerun()


# ============================================================
# DETALHE
# ============================================================

elif st.session_state.pagina == "Detalhe":
    fiscalizacao_id = st.session_state["fiscalizacao_selecionada"]

    if fiscalizacao_id is None:
        st.warning(t(t("Nenhuma fiscalização selecionada.")))
    else:
        try:
            detalhes = obter_detalhes_fiscalizacao(fiscalizacao_id)
        except Exception as erro:
            tratar_erro_operacional(erro, "abertura da fiscalização")
            detalhes = None

        if not detalhes:
            st.error(t(t("Fiscalização não encontrada.")))
        else:
            fiscalizacao = detalhes["fiscalizacao"]
            tributos = detalhes["tributos"]
            areas = detalhes["areas"]
            historico = detalhes["historico"]
            prazos = detalhes["prazos"]
            responsavel_principal = detalhes["responsavel_principal"]
            tarefa_estracta = detalhes.get("tarefa_estracta") or ""
            flash_executivo = detalhes.get("flash_report_executivo") or {
                "incluir_flash_report": False,
                "prioridade_executiva": "Média",
                "atualizacao_executiva": "",
            }
            tarefas = detalhes["tarefas"]
            documentos = detalhes.get("documentos") or []
            arquivamento = detalhes["arquivamento"]
            esta_arquivada = arquivamento["arquivada"]

            (
                id_,
                codigo,
                titulo,
                objeto,
                empresa,
                filial,
                orgao,
                exposicao,
                probabilidade,
                fase,
                data_entrada_fase,
                criado_em,
                atualizado_em,
            ) = fiscalizacao

            if (
                st.session_state.get("case_param_processado")
                != codigo
            ):
                st.session_state["case_param_processado"] = codigo
                definir_case_query_param(codigo)

            st.write("")

            voltar_coluna, _ = st.columns([1.8, 4.2])

            with voltar_coluna:
                st.button(
                    t("← Voltar para fiscalizações"),
                    key=f"voltar_fiscalizacoes_{id_}",
                    on_click=voltar_para_fiscalizacoes,
                    use_container_width=True,
                )

            st.caption("SYNGENTA BRASIL · TAX CONTROVERSY" if idioma_atual() == "pt" else "SYNGENTA BRAZIL · TAX CONTROVERSY")
            st.title(titulo)
            st.caption(f"{codigo} · {empresa}" + (f" · {filial}" if filial else ""))

            link_direto_fiscalizacao = montar_link_fiscalizacao(
                codigo
            )

            with st.container(border=True):
                lk1, lk2 = st.columns([4, 1])

                with lk1:
                    st.caption("LINK DIRETO DO CASO" if idioma_atual() == "pt" else "DIRECT CASE LINK")
                    st.code(
                        link_direto_fiscalizacao,
                        language=None,
                    )

                with lk2:
                    st.caption("ACESSO")
                    st.link_button(
                        "Abrir link",
                        link_direto_fiscalizacao,
                        use_container_width=True,
                    )

            if (
                obter_url_base_aplicacao()
                == BASE_URL_PADRAO
            ):
                st.caption(
                    (
                        "O link ainda usa localhost. Na publicação para a equipe, "
                        "configure FISCAL_TRACKER_BASE_URL com o endereço interno."
                    )
                )

            status1, status2, status3 = st.columns(3)
            status1.info(t("ARQUIVADA") if esta_arquivada else t("ATIVA"))
            status2.info(traduzir_valor(fase))
            if probabilidade == "Provável":
                status3.error(traduzir_valor(probabilidade))
            elif probabilidade == "Possível":
                status3.warning(traduzir_valor(probabilidade))
            else:
                status3.success(traduzir_valor(probabilidade))

            if flash_executivo.get("incluir_flash_report"):
                with st.container(border=True):
                    fx1, fx2 = st.columns([1, 3])
                    with fx1:
                        st.caption(
                            "FLASH REPORT EXECUTIVO"
                            if idioma_atual() == "pt"
                            else "EXECUTIVE FLASH REPORT"
                        )
                        st.success(
                            "Incluída"
                            if idioma_atual() == "pt"
                            else "Included"
                        )
                        st.caption(
                            (
                                "Prioridade: "
                                if idioma_atual() == "pt"
                                else "Priority: "
                            )
                            + traduzir_valor(
                                flash_executivo.get(
                                    "prioridade_executiva",
                                    "Média",
                                )
                            )
                        )
                    with fx2:
                        st.caption(
                            "ATUALIZAÇÃO EXECUTIVA"
                            if idioma_atual() == "pt"
                            else "EXECUTIVE UPDATE"
                        )
                        st.write(
                            flash_executivo.get(
                                "atualizacao_executiva",
                                "",
                            )
                            or (
                                "Ainda não informada. O Resumo Executivo do Report Fiscalizações "
                                "será usado como fallback no Dashboard."
                                if idioma_atual() == "pt"
                                else (
                                    "Not yet provided. The Tax Inspection Report Executive Summary "
                                    "will be used as fallback on the Dashboard."
                                )
                            )
                        )

            st.divider()

            with st.expander("Governança da fiscalização"):
                if esta_arquivada:
                    st.info("Esta fiscalização está arquivada.")
                    if arquivamento["arquivada_em"]:
                        st.caption(
                            "Arquivada em "
                            + formatar_data_hora(arquivamento["arquivada_em"])
                        )
                    if st.button(
                        "Reativar fiscalização",
                        type="primary",
                        key=f"reativar_{id_}",
                        use_container_width=True,
                    ):
                        try:
                            reativar_fiscalizacao(id_)
                            limpar_pdf_sessao(id_)
                            exibir_sucesso("Fiscalização reativada.")
                            st.rerun()
                        except Exception as erro:
                            tratar_erro_operacional(erro, "reativação da fiscalização")
                else:
                    st.warning(
                        "Arquivar retira o caso das visões operacionais, "
                        "mas preserva todo o histórico."
                    )
                    confirmar = st.checkbox(
                        "Confirmo o arquivamento.",
                        key=f"confirmar_arquivamento_{id_}",
                    )
                    if st.button(
                        "Arquivar fiscalização",
                        disabled=not confirmar,
                        key=f"arquivar_{id_}",
                        use_container_width=True,
                    ):
                        try:
                            arquivar_fiscalizacao(id_)
                            limpar_pdf_sessao(id_)
                            exibir_sucesso("Fiscalização arquivada.")
                            st.rerun()
                        except Exception as erro:
                            tratar_erro_operacional(erro, "arquivamento da fiscalização")

            if st.session_state.get("foco_detalhe") == "documentos":
                with st.container(border=True):
                    titulo_secao(
                        (
                            "Acesso direto — Documentos & Evidências"
                            if idioma_atual() == "pt"
                            else "Direct access — Documents & Evidence"
                        ),
                        (
                            "Acesso rápido aos anexos desta fiscalização a partir do Flash Report Executivo."
                            if idioma_atual() == "pt"
                            else "Quick access to this tax inspection's attachments from the Executive Flash Report."
                        ),
                    )
                    try:
                        documentos_rapidos = listar_documentos_fiscalizacao(id_)
                    except Exception as erro:
                        documentos_rapidos = []
                        tratar_erro_operacional(erro, "carregamento dos documentos")

                    if not documentos_rapidos:
                        st.info(
                            "Nenhum documento ou evidência foi adicionado a esta fiscalização."
                            if idioma_atual() == "pt"
                            else "No document or evidence has been added to this tax inspection."
                        )
                    else:
                        st.caption(
                            (f"{len(documentos_rapidos)} documento(s) disponível(is).")
                            if idioma_atual() == "pt"
                            else (f"{len(documentos_rapidos)} document(s) available.")
                        )
                        for documento in documentos_rapidos:
                            qd1, qd2 = st.columns([5, 1.5])
                            with qd1:
                                st.markdown(f"**{documento['nome_original']}**")
                                st.caption(
                                    f"{documento['tipo_documento']} · "
                                    f"{formatar_tamanho_arquivo(documento['tamanho_bytes'])}"
                                )
                            with qd2:
                                try:
                                    integridade_quick = verificar_integridade_documento(documento["id"])
                                    if integridade_quick.get("ok"):
                                        conteudo_quick = ler_conteudo_documento(documento["id"])
                                        st.download_button(
                                            ("Baixar" if idioma_atual() == "pt" else "Download"),
                                            data=conteudo_quick,
                                            file_name=documento["nome_original"],
                                            mime=documento["mime_type"] or "application/octet-stream",
                                            key=f"quick_download_doc_{documento['id']}",
                                            use_container_width=True,
                                        )
                                    else:
                                        st.error(
                                            integridade_quick.get("motivo") or
                                            ("Integridade não confirmada." if idioma_atual() == "pt" else "Integrity not confirmed.")
                                        )
                                except Exception as erro:
                                    tratar_erro_operacional(erro, "leitura do documento")

                    if st.button(
                        (
                            "Continuar na fiscalização completa"
                            if idioma_atual() == "pt"
                            else "Continue to full tax inspection"
                        ),
                        key=f"sair_foco_documentos_{id_}",
                    ):
                        st.session_state["foco_detalhe"] = ""
                        st.session_state["case_param_processado"] = f"{codigo}|"
                        definir_case_query_param(codigo)
                        st.rerun()

            aba_detalhe, aba_comunicacoes, aba_flash = st.tabs(
                [
                    "Detalhe",
                    t("Comunicações"),
                    t("Report Fiscalizações"),
                ]
            )

            with aba_detalhe:
                titulo_secao(t("Executive Overview"))
                status_prazo = calcular_status_prazo(
                    prazos["prazo_resposta"] if prazos else None,
                    fase,
                )

                e1, e2, e3, e4 = st.columns(4)
                e1.metric(t("Exposição"), formatar_milhoes(exposicao))
                e2.metric(t("Probabilidade"), probabilidade)
                e3.metric(t("Fase atual"), fase)
                e4.metric(
                    t("Prazo"),
                    formatar_data(prazos["prazo_resposta"] if prazos else None),
                )

                st.write("")
                titulo_secao(t("Case Information"))

                if esta_arquivada:
                    st.info(
                        "A edição fica bloqueada enquanto o caso estiver arquivado."
                    )
                else:
                    with st.expander(t("Editar dados cadastrais")):
                        areas_atuais = [area[0] for area in areas]
                        indice_probabilidade = (
                            PROBABILIDADES.index(probabilidade)
                            if probabilidade in PROBABILIDADES
                            else 1
                        )

                        with st.form(f"editar_fiscalizacao_{id_}"):
                            edit_titulo = st.text_input("Título *", value=titulo or "")
                            edit_objeto = st.text_area(
                                "Objeto / escopo",
                                value=objeto or "",
                                height=120,
                            )
                            ec1, ec2 = st.columns(2)
                            with ec1:
                                edit_empresa = st.text_input("Empresa *", value=empresa or "")
                            with ec2:
                                edit_filial = st.text_input("Filial", value=filial or "")
                            edit_responsavel = st.text_input(
                                "Responsável principal",
                                value=responsavel_principal or "",
                            )
                            edit_tarefa_estracta = st.text_input(
                                "Tarefa e-Stracta",
                                value=tarefa_estracta or "",
                                help=(
                                    "Informe o número ou código da tarefa correspondente "
                                    "na ferramenta interna e-Stracta."
                                ),
                            )
                            ec3, ec4 = st.columns(2)
                            with ec3:
                                edit_orgao = st.text_input(
                                    "Órgão fiscalizador",
                                    value=orgao or "",
                                )
                            with ec4:
                                edit_probabilidade = st.selectbox(
                                    t("Probabilidade"),
                                    PROBABILIDADES,
                                    index=indice_probabilidade,
                                    format_func=traduzir_opcao,
                                )
                            edit_exposicao = st.number_input(
                                "Exposição estimada (R$)",
                                min_value=0.0,
                                value=float(exposicao or 0),
                                step=1000.0,
                            )
                            edit_tributos = st.multiselect(
                                "Tributos envolvidos",
                                TRIBUTOS_DISPONIVEIS,
                                default=[
                                    item for item in tributos if item in TRIBUTOS_DISPONIVEIS
                                ],
                            )
                            areas_disponiveis_edicao = (
                                obter_nomes_areas_ativas()
                            )

                            # Preserva áreas antigas/inativas já vinculadas
                            # para que uma simples edição não apague o histórico.
                            for area_atual in areas_atuais:
                                if (
                                    area_atual
                                    not in areas_disponiveis_edicao
                                ):
                                    areas_disponiveis_edicao.append(
                                        area_atual
                                    )

                            edit_areas = st.multiselect(
                                "Áreas envolvidas",
                                areas_disponiveis_edicao,
                                default=areas_atuais,
                            )

                            st.divider()
                            st.subheader(
                                "Flash Report Executivo"
                                if idioma_atual() == "pt"
                                else "Executive Flash Report"
                            )
                            st.caption(
                                (
                                    "Marque esta opção quando a fiscalização tiver relevância "
                                    "executiva e deva aparecer no Flash Report do Dashboard."
                                )
                                if idioma_atual() == "pt"
                                else (
                                    "Select this option when the tax inspection is executive-relevant "
                                    "and should appear in the Dashboard Flash Report."
                                )
                            )
                            edit_incluir_flash = st.checkbox(
                                (
                                    "Incluir esta fiscalização no Flash Report"
                                    if idioma_atual() == "pt"
                                    else "Include this tax inspection in the Flash Report"
                                ),
                                value=bool(
                                    flash_executivo.get(
                                        "incluir_flash_report",
                                        False,
                                    )
                                ),
                            )
                            prioridade_atual = flash_executivo.get(
                                "prioridade_executiva",
                                "Média",
                            )
                            indice_prioridade = (
                                PRIORIDADES_EXECUTIVAS.index(prioridade_atual)
                                if prioridade_atual in PRIORIDADES_EXECUTIVAS
                                else 1
                            )
                            edit_prioridade_executiva = st.selectbox(
                                (
                                    "Prioridade executiva"
                                    if idioma_atual() == "pt"
                                    else "Executive priority"
                                ),
                                PRIORIDADES_EXECUTIVAS,
                                index=indice_prioridade,
                                format_func=traduzir_opcao,
                                help=(
                                    "Usada para ordenar os casos relevantes no Flash Report."
                                    if idioma_atual() == "pt"
                                    else "Used to order relevant cases in the Flash Report."
                                ),
                            )
                            edit_atualizacao_executiva = st.text_area(
                                (
                                    "Atualização executiva"
                                    if idioma_atual() == "pt"
                                    else "Executive update"
                                ),
                                value=flash_executivo.get(
                                    "atualizacao_executiva",
                                    "",
                                ),
                                height=100,
                                max_chars=500,
                                help=(
                                    "Resumo curto para o Dashboard. Se ficar vazio, o sistema poderá "
                                    "usar o Resumo Executivo do Report Fiscalizações."
                                    if idioma_atual() == "pt"
                                    else (
                                        "Short Dashboard update. If left blank, the system may use "
                                        "the Executive Summary from the Tax Inspection Report."
                                    )
                                ),
                            )

                            salvar = st.form_submit_button(
                                t("Salvar alterações"),
                                type="primary",
                                use_container_width=True,
                            )

                        if salvar:
                            erros = validar_campos_obrigatorios(
                                [
                                    ("o título da fiscalização", edit_titulo),
                                    ("a empresa", edit_empresa),
                                ]
                            )
                            if erros:
                                for erro in erros:
                                    st.warning(erro)
                            else:
                                try:
                                    atualizar_dados_fiscalizacao(
                                        fiscalizacao_id=id_,
                                        titulo=edit_titulo,
                                        objeto=edit_objeto,
                                        empresa=edit_empresa,
                                        filial=edit_filial,
                                        orgao=edit_orgao,
                                        exposicao=edit_exposicao,
                                        probabilidade=edit_probabilidade,
                                        tributos=edit_tributos,
                                        areas=edit_areas,
                                        responsavel_principal=edit_responsavel,
                                        tarefa_estracta=edit_tarefa_estracta,
                                        incluir_flash_report=edit_incluir_flash,
                                        prioridade_executiva=edit_prioridade_executiva,
                                        atualizacao_executiva=edit_atualizacao_executiva,
                                    )
                                    limpar_pdf_sessao(id_)
                                    exibir_sucesso(
                                        "Dados da fiscalização atualizados com sucesso."
                                    )
                                    st.rerun()
                                except Exception as erro:
                                    tratar_erro_operacional(
                                        erro,
                                        "atualização da fiscalização",
                                    )

                info1, info2 = st.columns(2)
                with info1:
                    with st.container(border=True):
                        st.caption("RESPONSÁVEL PRINCIPAL")
                        st.write(responsavel_principal or "-")
                        st.caption("ÓRGÃO FISCALIZADOR")
                        st.write(orgao or "-")
                        st.caption("TAREFA E-STRACTA")
                        st.write(tarefa_estracta or "-")
                with info2:
                    with st.container(border=True):
                        st.caption("OBJETO / ESCOPO")
                        st.write(objeto or "-")
                        st.caption("TRIBUTOS")
                        st.write(", ".join(tributos) if tributos else "-")
                        st.caption("DOCUMENTOS NO DOSSIÊ")
                        st.write(len(documentos))

                st.write("")
                titulo_secao(t("Deadlines"))
                pd1, pd2, pd3 = st.columns(3)
                pd1.metric(
                    t("Recebimento"),
                    formatar_data(prazos["data_recebimento"] if prazos else None),
                )
                pd2.metric(
                    t("Prazo atual"),
                    formatar_data(prazos["prazo_resposta"] if prazos else None),
                )
                with pd3:
                    exibir_status_prazo(
                        status_prazo["status"],
                        status_prazo["dias_restantes"],
                    )

                if not esta_arquivada:
                    with st.expander(t("Editar prazo")):
                        data_recebimento = prazos["data_recebimento"] if prazos else None
                        prazo_resposta = prazos["prazo_resposta"] if prazos else None
                        recebimento_padrao = (
                            date.fromisoformat(data_recebimento)
                            if data_recebimento
                            else date.today()
                        )
                        prazo_padrao = (
                            date.fromisoformat(prazo_resposta)
                            if prazo_resposta
                            else date.today()
                        )

                        with st.form(f"prazo_{id_}"):
                            pp1, pp2 = st.columns(2)
                            with pp1:
                                nova_data = st.date_input(
                                    "Data de recebimento",
                                    value=recebimento_padrao,
                                    format="DD/MM/YYYY",
                                )
                            with pp2:
                                novo_prazo = st.date_input(
                                    "Prazo de resposta",
                                    value=prazo_padrao,
                                    format="DD/MM/YYYY",
                                )
                            salvar_prazo = st.form_submit_button(
                                t("Salvar prazo"),
                                type="primary",
                                use_container_width=True,
                            )

                        if salvar_prazo:
                            erro_periodo = validar_periodo(
                                nova_data,
                                novo_prazo,
                                "data de recebimento",
                                "data de resposta",
                            )
                            if erro_periodo:
                                st.warning(erro_periodo)
                            else:
                                try:
                                    atualizar_prazos_fiscalizacao(
                                        id_,
                                        nova_data,
                                        novo_prazo,
                                    )
                                    limpar_pdf_sessao(id_)
                                    exibir_sucesso(t(t("Prazo atualizado com sucesso.")))
                                    st.rerun()
                                except Exception as erro:
                                    tratar_erro_operacional(erro, "atualização do prazo")

                st.write("")
                titulo_secao(t("Internal Tasks"))

                if not esta_arquivada:
                    with st.expander(t("Nova tarefa")):
                        with st.form(f"nova_tarefa_{id_}", clear_on_submit=True):
                            tarefa_titulo = st.text_input("Título *")
                            tarefa_descricao = st.text_area("Descrição")
                            ta1, ta2 = st.columns(2)
                            with ta1:
                                tarefa_responsavel = st.text_input("Responsável")
                            with ta2:
                                tarefa_prazo = st.date_input(
                                    "Prazo interno",
                                    value=date.today(),
                                    format="DD/MM/YYYY",
                                )
                            criar = st.form_submit_button(
                                t("Criar tarefa"),
                                type="primary",
                                use_container_width=True,
                            )

                        if criar:
                            erros = validar_campos_obrigatorios(
                                [("o título da tarefa", tarefa_titulo)]
                            )
                            if erros:
                                for erro in erros:
                                    st.warning(erro)
                            else:
                                try:
                                    cadastrar_tarefa(
                                        fiscalizacao_id=id_,
                                        titulo=tarefa_titulo.strip(),
                                        descricao=tarefa_descricao,
                                        responsavel=tarefa_responsavel,
                                        prazo=tarefa_prazo,
                                    )
                                    limpar_pdf_sessao(id_)
                                    exibir_sucesso(t(t("Tarefa criada com sucesso.")))
                                    st.rerun()
                                except Exception as erro:
                                    tratar_erro_operacional(erro, "criação da tarefa")

                if not tarefas:
                    st.info(t(t("Nenhuma tarefa cadastrada.")))
                else:
                    for tarefa in tarefas:
                        (
                            tarefa_id,
                            tarefa_titulo,
                            tarefa_descricao,
                            tarefa_responsavel,
                            tarefa_prazo,
                            tarefa_status,
                            tarefa_concluida_em,
                            tarefa_criada_em,
                            tarefa_atualizada_em,
                        ) = tarefa

                        prazo_tarefa = calcular_status_prazo_tarefa(
                            tarefa_prazo,
                            tarefa_status,
                        )

                        with st.container(border=True):
                            tr1, tr2, tr3 = st.columns([4, 2, 2])
                            with tr1:
                                st.markdown(f"**{tarefa_titulo}**")
                                if tarefa_descricao:
                                    st.write(tarefa_descricao)
                                st.caption(
                                    f"Responsável: {tarefa_responsavel or '-'}"
                                )
                            with tr2:
                                st.caption("PRAZO")
                                st.write(formatar_data(tarefa_prazo))
                                st.caption(traduzir_valor(prazo_tarefa["status_prazo"]))
                            with tr3:
                                st.caption("STATUS")
                                if esta_arquivada:
                                    st.write(traduzir_valor(tarefa_status))
                                else:
                                    novo_status = st.selectbox(
                                        "Status",
                                        STATUS_TAREFAS,
                                        index=STATUS_TAREFAS.index(tarefa_status),
                                        format_func=traduzir_opcao,
                                        key=f"status_{tarefa_id}",
                                        label_visibility="collapsed",
                                    )
                                    if st.button(
                                        t("Atualizar"),
                                        key=f"update_{tarefa_id}",
                                        use_container_width=True,
                                    ):
                                        if novo_status == tarefa_status:
                                            st.info(
                                                "Nenhuma alteração de status foi realizada."
                                            )
                                        else:
                                            try:
                                                atualizar_status_tarefa(
                                                    tarefa_id,
                                                    novo_status,
                                                )
                                                limpar_pdf_sessao(id_)
                                                exibir_sucesso(
                                                    "Status da tarefa atualizado."
                                                )
                                                st.rerun()
                                            except Exception as erro:
                                                tratar_erro_operacional(
                                                    erro,
                                                    "atualização da tarefa",
                                                )

                st.write("")
                titulo_secao(
                    t("Governança do encerramento"),
                    (
                        "Defina os aprovadores responsáveis pela ciência final da fiscalização. "
                        "O encerramento só é liberado quando existir pelo menos um aprovador "
                        "essencial e todos os essenciais tiverem registrado ciência."
                    ),
                )

                try:
                    candidatos_encerramento = listar_candidatos_aprovacao_encerramento(id_)
                    aprovacoes_encerramento = listar_aprovacoes_encerramento(id_)
                    status_encerramento = obter_status_encerramento(id_)
                except Exception as erro:
                    candidatos_encerramento = []
                    aprovacoes_encerramento = []
                    status_encerramento = {
                        "total_aprovadores": 0,
                        "total_essenciais": 0,
                        "essenciais_cientes": 0,
                        "essenciais_pendentes": 0,
                        "pode_encerrar": False,
                        "pendentes": [],
                    }
                    tratar_erro_operacional(
                        erro,
                        "consulta da governança de encerramento",
                    )

                ge1, ge2, ge3 = st.columns(3)
                ge1.metric(
                    t("Aprovadores cadastrados"),
                    status_encerramento.get("total_aprovadores", 0),
                )
                ge2.metric(
                    t("Aprovadores essenciais"),
                    status_encerramento.get("total_essenciais", 0),
                )
                ge3.metric(
                    t("Essenciais pendentes"),
                    status_encerramento.get("essenciais_pendentes", 0),
                )

                if fase == "Encerrada":
                    st.success(
                        "Fiscalização encerrada com governança de aprovação concluída."
                    )
                elif esta_arquivada:
                    st.info(
                        "A gestão de aprovadores fica indisponível enquanto a fiscalização estiver arquivada."
                    )
                else:
                    if status_encerramento.get("pode_encerrar"):
                        st.success(
                            "Todos os aprovadores essenciais registraram ciência. O encerramento está liberado."
                        )
                    else:
                        if status_encerramento.get("total_essenciais", 0) == 0:
                            st.warning(
                                "O encerramento está bloqueado: defina pelo menos um aprovador essencial."
                            )
                        elif status_encerramento.get("essenciais_pendentes", 0) > 0:
                            nomes_pendentes = [
                                f"{item.get('nome', '-')} ({item.get('area', '-')})"
                                for item in status_encerramento.get("pendentes", [])
                            ]
                            st.warning(
                                "O encerramento está bloqueado. Aguardando ciência de: "
                                + ", ".join(nomes_pendentes)
                            )

                    with st.expander("Adicionar ou atualizar aprovador"):
                        if not candidatos_encerramento:
                            st.info(
                                "Nenhum candidato disponível. Vincule pessoas ativas às áreas envolvidas nesta fiscalização."
                            )
                        else:
                            mapa_candidatos = {
                                (
                                    f"{item['area']} · {item['nome']}"
                                    + (f" · {item['papel']}" if item.get('papel') else "")
                                ): item
                                for item in candidatos_encerramento
                            }

                            with st.form(f"novo_aprovador_encerramento_{id_}"):
                                candidato_label = st.selectbox(
                                    "Pessoa / área *",
                                    list(mapa_candidatos.keys()),
                                )
                                candidato = mapa_candidatos[candidato_label]
                                papel_aprovador = st.text_input(
                                    "Papel no encerramento",
                                    value=(candidato.get("papel") or "Aprovador"),
                                    help="Ex.: Reporting, Advisory, Tax Controversy.",
                                )
                                aprovador_essencial = st.checkbox(
                                    "Aprovador essencial",
                                    value=False,
                                    help=(
                                        "A ciência deste aprovador será obrigatória para permitir o encerramento."
                                    ),
                                )
                                salvar_aprovador = st.form_submit_button(
                                    t("Salvar aprovador"),
                                    type="primary",
                                    use_container_width=True,
                                )

                            if salvar_aprovador:
                                try:
                                    definir_aprovador_encerramento(
                                        fiscalizacao_id=id_,
                                        pessoa_id=candidato["pessoa_id"],
                                        area_id=candidato["area_id"],
                                        essencial=aprovador_essencial,
                                        papel=papel_aprovador,
                                    )
                                    exibir_sucesso(
                                        "Aprovador de encerramento salvo com sucesso."
                                    )
                                    st.rerun()
                                except Exception as erro:
                                    tratar_erro_operacional(
                                        erro,
                                        "cadastro do aprovador de encerramento",
                                    )

                if not aprovacoes_encerramento:
                    st.info("Nenhum aprovador de encerramento cadastrado.")
                else:
                    for aprovador in aprovacoes_encerramento:
                        with st.container(border=True):
                            ap1, ap2, ap3 = st.columns([4, 2, 2])

                            with ap1:
                                titulo_aprovador = (
                                    f"✓ {aprovador['nome']}"
                                    if aprovador["status"] == "Ciente"
                                    else f"○ {aprovador['nome']}"
                                )
                                st.markdown(f"**{titulo_aprovador}**")
                                st.write(aprovador.get("area") or "-")
                                if aprovador.get("papel"):
                                    st.caption(aprovador["papel"])
                                if aprovador.get("email"):
                                    st.caption(aprovador["email"])

                            with ap2:
                                st.caption("CLASSIFICAÇÃO")
                                if aprovador["essencial"]:
                                    st.error(t("Essencial"))
                                else:
                                    st.info(t("Aprovador"))

                                st.caption("STATUS")
                                if aprovador["status"] == "Ciente":
                                    st.success(t("Ciente"))
                                    if aprovador.get("ciente_em"):
                                        st.caption(
                                            formatar_data_hora(
                                                aprovador["ciente_em"]
                                            )
                                        )
                                else:
                                    st.warning(t("Pendente"))

                            with ap3:
                                if fase != "Encerrada" and not esta_arquivada:
                                    if aprovador["status"] == "Ciente":
                                        motivo_revogacao = st.text_input(
                                            "Motivo da revogação",
                                            key=f"motivo_revogacao_{aprovador['id']}",
                                            placeholder="Opcional",
                                        )
                                        if st.button(
                                            "Revogar ciência",
                                            key=f"revogar_ciencia_{aprovador['id']}",
                                            use_container_width=True,
                                        ):
                                            try:
                                                revogar_ciencia_encerramento(
                                                    aprovador["id"],
                                                    motivo_revogacao,
                                                )
                                                exibir_sucesso(
                                                    "Ciência revogada com sucesso."
                                                )
                                                st.rerun()
                                            except Exception as erro:
                                                tratar_erro_operacional(
                                                    erro,
                                                    "revogação da ciência",
                                                )
                                    else:
                                        observacao_ciencia = st.text_input(
                                            "Observação da ciência",
                                            key=f"obs_ciencia_{aprovador['id']}",
                                            placeholder="Opcional",
                                        )
                                        if st.button(
                                            "Registrar ciência",
                                            key=f"registrar_ciencia_{aprovador['id']}",
                                            type="primary" if aprovador["essencial"] else "secondary",
                                            use_container_width=True,
                                        ):
                                            try:
                                                registrar_ciencia_encerramento(
                                                    aprovador["id"],
                                                    observacao_ciencia,
                                                )
                                                exibir_sucesso(
                                                    "Ciência registrada com sucesso."
                                                )
                                                st.rerun()
                                            except Exception as erro:
                                                tratar_erro_operacional(
                                                    erro,
                                                    "registro da ciência",
                                                )

                                        if st.button(
                                            "Remover aprovador",
                                            key=f"remover_aprovador_{aprovador['id']}",
                                            use_container_width=True,
                                        ):
                                            try:
                                                remover_aprovador_encerramento(
                                                    aprovador["id"]
                                                )
                                                exibir_sucesso(
                                                    "Aprovador removido com sucesso."
                                                )
                                                st.rerun()
                                            except Exception as erro:
                                                tratar_erro_operacional(
                                                    erro,
                                                    "remoção do aprovador",
                                                )

                st.write("")
                titulo_secao(t("Workflow"))

                if esta_arquivada:
                    st.info(t(t("Workflow indisponível para registros arquivados.")))
                elif fase == "Encerrada":
                    st.success(t(t("Fiscalização encerrada.")))
                else:
                    proxima_fase = calcular_proxima_fase(fase)
                    with st.container(border=True):
                        wf1, wf2 = st.columns([2, 3])
                        with wf1:
                            st.caption("FASE ATUAL")
                            st.write(traduzir_valor(fase))
                            st.caption("PRÓXIMA FASE SUGERIDA")
                            st.write(traduzir_valor(proxima_fase))
                            st.progress(
                                (FASES_FISCALIZACAO.index(fase) + 1)
                                / len(FASES_FISCALIZACAO)
                            )
                        with wf2:
                            with st.form(f"workflow_{id_}"):
                                nova_fase = st.selectbox(
                                    t("Nova fase"),
                                    FASES_FISCALIZACAO,
                                    index=(
                                        FASES_FISCALIZACAO.index(proxima_fase)
                                        if proxima_fase
                                        else FASES_FISCALIZACAO.index(fase)
                                    ),
                                    format_func=traduzir_opcao,
                                )
                                observacao = st.text_area(t("Observação"))
                                mover = st.form_submit_button(
                                    t("Atualizar fase"),
                                    type="primary",
                                    use_container_width=True,
                                )

                            if mover:
                                if nova_fase == fase:
                                    st.info(t(t("A fiscalização já está nesta fase.")))
                                else:
                                    try:
                                        observacao_workflow = observacao
                                        if nova_fase == "Encerrada" and not observacao_workflow.strip():
                                            observacao_workflow = (
                                                "Fiscalização encerrada após conclusão do fluxo de ciências obrigatórias."
                                            )
                                        alterar_fase_fiscalizacao(
                                            id_,
                                            nova_fase,
                                            observacao_workflow,
                                        )
                                        limpar_pdf_sessao(id_)
                                        if nova_fase == "Encerrada":
                                            st.session_state[f"encerramento_concluido_{id_}"] = True
                                            st.session_state[f"abrir_comunicacao_encerramento_{id_}"] = True
                                            exibir_sucesso(
                                                "Fiscalização encerrada após conclusão do fluxo de ciências obrigatórias."
                                            )
                                        else:
                                            exibir_sucesso(t(t("Fase atualizada com sucesso.")))
                                        st.rerun()
                                    except Exception as erro:
                                        tratar_erro_operacional(
                                            erro,
                                            "movimentação do workflow",
                                        )

                st.write("")
                titulo_secao(t("Stakeholders & Areas"))

                if not areas:
                    st.info("Nenhuma área cadastrada.")
                else:
                    colunas_area = st.columns(min(3, len(areas)))
                    for indice, area in enumerate(areas):
                        nome_area, tipo, status_area = area
                        coluna = colunas_area[indice % len(colunas_area)]
                        with coluna:
                            with st.container(border=True):
                                st.markdown(f"**{nome_area}**")
                                st.write(tipo)
                                st.caption(traduzir_valor(status_area))

                st.write("")
                titulo_secao(
                    t("Documents & Evidence"),
                    ("Documentos, termos, e-mails, planilhas e demais evidências relacionadas à fiscalização." if idioma_atual() == "pt" else "Documents, notices, emails, spreadsheets and other evidence related to the tax inspection."),
                )

                try:
                    documentos = listar_documentos_fiscalizacao(id_)
                except Exception as erro:
                    documentos = []
                    tratar_erro_operacional(erro, "carregamento dos documentos")

                if esta_arquivada:
                    st.info(
                        "A fiscalização está arquivada. Os documentos permanecem "
                        "disponíveis para consulta e download, mas novos arquivos não podem ser adicionados."
                    )
                else:
                    with st.expander("Adicionar documento ou evidência"):
                        with st.form(f"novo_documento_{id_}", clear_on_submit=True):
                            doc_tipo = st.selectbox("Tipo do documento *", TIPOS_DOCUMENTO)
                            doc_descricao = st.text_area(
                                "Descrição",
                                placeholder=(
                                    "Descreva brevemente o conteúdo ou relevância do documento."
                                ),
                                height=90,
                            )
                            doc_data = st.date_input(
                                "Data do documento",
                                value=date.today(),
                                format="DD/MM/YYYY",
                            )
                            doc_arquivo = st.file_uploader(
                                "Arquivo *",
                                type=EXTENSOES_UPLOAD,
                                help=(
                                    "Formatos permitidos: PDF, Word, Excel, CSV, TXT, MSG, "
                                    "EML e imagens. Limite: 25 MB por arquivo."
                                ),
                            )
                            adicionar_doc = st.form_submit_button(
                                t("Adicionar ao dossiê"),
                                type="primary",
                                use_container_width=True,
                            )

                        if adicionar_doc:
                            if doc_arquivo is None:
                                st.warning("Selecione um arquivo.")
                            else:
                                try:
                                    adicionar_documento_fiscalizacao(
                                        fiscalizacao_id=id_,
                                        tipo_documento=doc_tipo,
                                        nome_original=doc_arquivo.name,
                                        conteudo=doc_arquivo.getvalue(),
                                        descricao=doc_descricao,
                                        data_documento=doc_data,
                                        mime_type=doc_arquivo.type or None,
                                    )
                                    limpar_pdf_sessao(id_)
                                    exibir_sucesso(
                                        "Documento adicionado ao dossiê com sucesso."
                                    )
                                    st.rerun()
                                except Exception as erro:
                                    tratar_erro_operacional(erro, "inclusão do documento")

                if not documentos:
                    st.info(
                        "Nenhum documento ou evidência foi adicionado a esta fiscalização."
                    )
                else:
                    st.caption(f"{len(documentos)} documento(s) no dossiê.")
                    for documento in documentos:
                        documento_id = documento["id"]
                        with st.container(border=True):
                            dc1, dc2, dc3 = st.columns([5, 2, 2])
                            with dc1:
                                st.caption(documento["tipo_documento"].upper())
                                st.markdown(f"**{documento['nome_original']}**")
                                if documento["descricao"]:
                                    st.write(documento["descricao"])
                                detalhes_documento = []
                                if documento["data_documento"]:
                                    detalhes_documento.append(
                                        "Documento: "
                                        + formatar_data(documento["data_documento"])
                                    )
                                detalhes_documento.append(
                                    "Incluído: " + formatar_data_hora(documento["criado_em"])
                                )
                                detalhes_documento.append(
                                    formatar_tamanho_arquivo(documento["tamanho_bytes"])
                                )
                                st.caption(" · ".join(detalhes_documento))
                                st.caption(
                                    f"SHA-256: {documento['hash_sha256'][:16]}…"
                                )
                            with dc2:
                                st.caption("INTEGRIDADE")
                                try:
                                    integridade = verificar_integridade_documento(documento_id)
                                    if integridade["ok"]:
                                        st.success("Arquivo íntegro")
                                    else:
                                        st.error(integridade["motivo"])
                                except Exception:
                                    integridade = {"ok": False}
                                    st.error("Não verificada")

                                if integridade.get("ok"):
                                    try:
                                        conteudo_documento = ler_conteudo_documento(
                                            documento_id
                                        )
                                        st.download_button(
                                            "Baixar",
                                            data=conteudo_documento,
                                            file_name=documento["nome_original"],
                                            mime=documento["mime_type"]
                                            or "application/octet-stream",
                                            key=f"download_doc_{documento_id}",
                                            use_container_width=True,
                                        )
                                    except Exception as erro:
                                        tratar_erro_operacional(erro, "leitura do documento")
                            with dc3:
                                st.caption("GESTÃO")
                                if esta_arquivada:
                                    st.caption("Somente consulta")
                                else:
                                    confirmar_exclusao = st.checkbox(
                                        t("Confirmar exclusão"),
                                        key=f"confirmar_doc_{documento_id}",
                                    )
                                    if st.button(
                                        t("Excluir"),
                                        key=f"excluir_doc_{documento_id}",
                                        disabled=not confirmar_exclusao,
                                        use_container_width=True,
                                    ):
                                        try:
                                            excluir_documento_fiscalizacao(documento_id)
                                            limpar_pdf_sessao(id_)
                                            exibir_sucesso("Documento excluído.")
                                            st.rerun()
                                        except Exception as erro:
                                            tratar_erro_operacional(erro, "exclusão do documento")

                st.write("")
                titulo_secao(
                    t("Timeline"),
                    ("Histórico dos principais desdobramentos da fiscalização." if idioma_atual() == "pt" else "History of the main developments in the tax inspection."),
                )

                if not historico:
                    st.info("Nenhum evento registrado.")
                else:
                    for item in historico:
                        historico_id, historico_fase, historico_data, historico_observacao, historico_criado_em = item
                        with st.container(border=True):
                            tl1, tl2 = st.columns([1, 5])
                            with tl1:
                                st.caption("DATA")
                                st.write(formatar_data(historico_data))
                            with tl2:
                                st.markdown(f"**{traduzir_valor(historico_fase)}**")
                                if historico_observacao:
                                    st.write(historico_observacao)
                                st.caption(
                                    "Registrado em "
                                    + formatar_data_hora(historico_criado_em)
                                )


            # =================================================
            # ABA COMUNICACOES
            # FASE 2.0B3
            # =================================================

            with aba_comunicacoes:

                titulo_secao(
                    t("Comunicações"),
                    (
                        ("Prepare comunicações para as áreas e pessoas envolvidas, abra o rascunho no Outlook e preserve o histórico no dossiê da fiscalização." if idioma_atual() == "pt" else "Prepare communications for involved areas and people, open the draft in Outlook and preserve the history in the tax inspection case file.")
                    ),
                )

                st.info(
                    (
                        ("O envio automático pelo Microsoft Graph depende de liberação do ambiente corporativo. Nesta etapa, o Fiscal Tracker prepara o e-mail e abre o Outlook. O envio continua sendo confirmado por você." if idioma_atual() == "pt" else "Automatic sending through Microsoft Graph depends on corporate environment approval. At this stage, Fiscal Tracker prepares the email and opens Outlook. You remain responsible for confirming the send action.")
                    )
                )

                try:

                    destinatarios_automaticos = (
                        obter_destinatarios_fiscalizacao(
                            id_
                        )
                    )

                except Exception as erro:

                    destinatarios_automaticos = []

                    tratar_erro_operacional(
                        erro,
                        "consulta dos destinatários da fiscalização",
                    )

                try:

                    historico_comunicacoes = (
                        listar_comunicacoes_fiscalizacao(
                            id_
                        )
                    )

                except Exception as erro:

                    historico_comunicacoes = []

                    tratar_erro_operacional(
                        erro,
                        "consulta do histórico de comunicações",
                    )

                # ---------------------------------------------
                # RESUMO
                # ---------------------------------------------

                total_enviadas = sum(
                    1
                    for comunicacao
                    in historico_comunicacoes
                    if comunicacao[
                        "status"
                    ]
                    == "Enviada"
                )

                total_preparadas = sum(
                    1
                    for comunicacao
                    in historico_comunicacoes
                    if comunicacao[
                        "status"
                    ]
                    in [
                        "Rascunho",
                        "Preparada",
                    ]
                )

                cm1, cm2, cm3 = (
                    st.columns(3)
                )

                with cm1:

                    st.metric(
                        t("Destinatários automáticos"),
                        len(
                            destinatarios_automaticos
                        ),
                    )

                with cm2:

                    st.metric(
                        t("Comunicações enviadas"),
                        total_enviadas,
                    )

                with cm3:

                    st.metric(
                        t("Pendentes de confirmação"),
                        total_preparadas,
                    )

                st.write("")

                # ---------------------------------------------
                # DESTINATARIOS ATUAIS
                # ---------------------------------------------

                titulo_secao(
                    t("Recipients"),
                    (
                        ("Os destinatários são obtidos das áreas vinculadas à fiscalização e dos envolvidos ativos cadastrados nessas áreas." if idioma_atual() == "pt" else "Recipients are obtained from areas linked to the tax inspection and active stakeholders registered in those areas.")
                    ),
                )

                if not destinatarios_automaticos:

                    st.warning(
                        (
                            "Nenhum destinatário automático foi encontrado. "
                            "Cadastre e-mails em Áreas & Envolvidos e confirme "
                            "que as áreas estão vinculadas a esta fiscalização."
                        )
                    )

                else:

                    for indice_destinatario, destinatario in enumerate(
                        destinatarios_automaticos,
                        start=1,
                    ):

                        with st.container(
                            border=True
                        ):

                            dr1, dr2 = (
                                st.columns(
                                    [4, 2]
                                )
                            )

                            with dr1:

                                nome_destinatario = (
                                    destinatario.get(
                                        "nome"
                                    )
                                    or (
                                        "Destinatário "
                                        f"{indice_destinatario}"
                                    )
                                )

                                st.markdown(
                                    f"**{nome_destinatario}**"
                                )

                                st.write(
                                    destinatario.get(
                                        "email"
                                    )
                                    or "-"
                                )

                            with dr2:

                                st.caption(
                                    "ORIGEM / PAPEL"
                                )

                                st.write(
                                    destinatario.get(
                                        "origem"
                                    )
                                    or "-"
                                )

                                st.caption(
                                    destinatario.get(
                                        "papel"
                                    )
                                    or "-"
                                )

                st.write("")

                # ---------------------------------------------
                # NOVA COMUNICACAO
                # ---------------------------------------------

                titulo_secao(
                    t("Prepare Communication")
                )

                if esta_arquivada:

                    st.info(
                        (
                            "A fiscalização está arquivada. "
                            "Reative o caso antes de preparar "
                            "uma nova comunicação."
                        )
                    )

                elif not destinatarios_automaticos:

                    st.info(
                        (
                            "A nova comunicação ficará disponível "
                            "quando houver pelo menos um destinatário "
                            "automático cadastrado para o caso."
                        )
                    )

                else:

                    prazo_comunicacao = (
                        prazos[
                            "prazo_resposta"
                        ]
                        if prazos
                        else None
                    )

                    modo_encerramento = bool(
                        fase == "Encerrada"
                        and st.session_state.get(
                            f"abrir_comunicacao_encerramento_{id_}", False
                        )
                    )

                    if modo_encerramento:
                        assunto_padrao = (
                            f"[{codigo}] Encerramento da fiscalização - {titulo}"
                            if idioma_atual() == "pt"
                            else f"[{codigo}] Tax inspection closure - {titulo}"
                        )
                        mensagem_padrao = montar_mensagem_encerramento_padrao(
                            codigo=codigo,
                            titulo=titulo,
                            empresa=empresa,
                            link_fiscalizacao=link_direto_fiscalizacao,
                        )
                    else:
                        assunto_padrao = (
                            f"[{codigo}] Atualização da fiscalização - {titulo}"
                            if idioma_atual() == "pt"
                            else f"[{codigo}] Tax inspection update - {titulo}"
                        )
                        mensagem_padrao = montar_mensagem_comunicacao_padrao(
                            codigo=codigo,
                            titulo=titulo,
                            empresa=empresa,
                            fase=fase,
                            prazo_resposta=prazo_comunicacao,
                            tarefa_estracta=tarefa_estracta,
                            link_fiscalizacao=link_direto_fiscalizacao,
                        )

                    opcoes_destinatarios = {
                        (
                            f"{item.get('nome') or item.get('email')} "
                            f"<{item.get('email')}>"
                        ): item
                        for item
                        in destinatarios_automaticos
                    }

                    with st.form(
                        f"comunicacao_{id_}"
                    ):

                        tipo_comunicacao = (
                            st.selectbox(
                                t("Tipo da comunicação"),
                                TIPOS_COMUNICACAO,
                                index=(
                                    TIPOS_COMUNICACAO.index("Encerramento")
                                    if modo_encerramento and "Encerramento" in TIPOS_COMUNICACAO
                                    else 0
                                ),
                                format_func=traduzir_opcao,
                            )
                        )

                        selecionados = (
                            st.multiselect(
                                t("Destinatários *"),
                                options=list(
                                    opcoes_destinatarios.keys()
                                ),
                                default=list(
                                    opcoes_destinatarios.keys()
                                ),
                            )
                        )

                        assunto_comunicacao = (
                            st.text_input(
                                t("Assunto *"),
                                value=(
                                    assunto_padrao
                                ),
                            )
                        )

                        mensagem_comunicacao = (
                            st.text_area(
                                t("Mensagem *"),
                                value=(
                                    mensagem_padrao
                                ),
                                height=260,
                            )
                        )

                        preparar_comunicacao = (
                            st.form_submit_button(
                                t("Preparar comunicação"),
                                type="primary",
                                use_container_width=True,
                            )
                        )

                    if preparar_comunicacao:

                        destinatarios_selecionados = [
                            opcoes_destinatarios[
                                chave
                            ]
                            for chave
                            in selecionados
                        ]

                        erros_comunicacao = (
                            validar_campos_obrigatorios(
                                [
                                    (
                                        "o assunto da comunicação",
                                        assunto_comunicacao,
                                    ),
                                    (
                                        "a mensagem da comunicação",
                                        mensagem_comunicacao,
                                    ),
                                ]
                            )
                        )

                        if not destinatarios_selecionados:

                            erros_comunicacao.append(
                                "Selecione pelo menos um destinatário."
                            )

                        if erros_comunicacao:

                            for erro in erros_comunicacao:

                                st.warning(
                                    erro
                                )

                        else:

                            try:

                                comunicacao_id = (
                                    criar_comunicacao(
                                        fiscalizacao_id=id_,
                                        assunto=(
                                            assunto_comunicacao
                                            .strip()
                                        ),
                                        mensagem=(
                                            mensagem_comunicacao
                                            .strip()
                                        ),
                                        destinatarios=(
                                            destinatarios_selecionados
                                        ),
                                        tipo=(
                                            tipo_comunicacao
                                        ),
                                        status="Preparada",
                                    )
                                )

                                st.session_state[
                                    f"comunicacao_preparada_{id_}"
                                ] = comunicacao_id

                                exibir_sucesso(
                                    (
                                        "Comunicação preparada e "
                                        "registrada no histórico."
                                    )
                                )

                                st.rerun()

                            except Exception as erro:

                                tratar_erro_operacional(
                                    erro,
                                    "preparação da comunicação",
                                )

                # ---------------------------------------------
                # ENCERRAMENTO GOVERNADO - FASE 2.0D4
                # ---------------------------------------------

                if fase == "Encerrada":
                    st.write("")
                    titulo_secao(
                        t("Comunicação final de encerramento"),
                        (
                            "Gere o Report Fiscalizações Final, anexe-o ao rascunho do Outlook "
                            "e preserve a confirmação do envio no histórico da fiscalização."
                        ),
                    )

                    try:
                        status_final = obter_status_encerramento(id_)
                    except Exception:
                        status_final = {}

                    ef1, ef2, ef3 = st.columns(3)
                    ef1.metric(t("Aprovadores essenciais"), status_final.get("total_essenciais", 0))
                    ef2.metric(t("Ciências essenciais"), status_final.get("essenciais_cientes", 0))
                    ef3.metric(t("Pendências"), status_final.get("essenciais_pendentes", 0))

                    chave_enc_pdf = f"encerramento_pdf_{id_}"
                    chave_enc_nome = f"encerramento_pdf_nome_{id_}"

                    with st.container(border=True):
                        st.markdown("**1. Report Fiscalizações Final**")
                        st.caption(
                            "O relatório final contém a trilha de aprovações e os anexos compatíveis consolidados no 2.0D3."
                        )
                        ec1, ec2 = st.columns(2)
                        with ec1:
                            if st.button(
                                t("Gerar Report Fiscalizações Final"),
                                key=f"gerar_flash_encerramento_{id_}",
                                type="primary",
                                use_container_width=True,
                            ):
                                try:
                                    pdf_final = gerar_flash_report_pdf(
                                        id_, incorporar_anexos=True
                                    )
                                    st.session_state[chave_enc_pdf] = pdf_final
                                    st.session_state[chave_enc_nome] = (
                                        f"Report_Fiscalizacao_Final_{codigo}.pdf"
                                    )
                                    exibir_sucesso(
                                        "Report Fiscalizações Final gerado e pronto para a comunicação final."
                                    )
                                except Exception as erro:
                                    tratar_erro_operacional(
                                        erro, "geração do Report Fiscalizações Final"
                                    )
                        with ec2:
                            if chave_enc_pdf in st.session_state:
                                st.download_button(
                                    t("Baixar Report Fiscalizações Final"),
                                    data=st.session_state[chave_enc_pdf],
                                    file_name=st.session_state[chave_enc_nome],
                                    mime="application/pdf",
                                    use_container_width=True,
                                )

                    with st.container(border=True):
                        st.markdown("**2. Rascunho de encerramento no Outlook**")
                        st.caption(
                            "O Fiscal Tracker prepara um arquivo .eml com destinatários, assunto, mensagem "
                            "e o Report Fiscalizações Final já anexado. Abra o arquivo no Outlook, revise e envie."
                        )
                        if chave_enc_pdf not in st.session_state:
                            st.warning(
                                "Gere o Report Fiscalizações Final antes de concluir a comunicação de encerramento."
                            )
                        else:
                            st.success(
                                f"Arquivo preparado: {st.session_state[chave_enc_nome]}"
                            )

                        if st.button(
                            "Usar modelo de encerramento",
                            key=f"usar_modelo_encerramento_{id_}",
                            use_container_width=True,
                        ):
                            st.session_state[f"abrir_comunicacao_encerramento_{id_}"] = True
                            st.rerun()

                # ---------------------------------------------
                # COMUNICACAO PREPARADA
                # ---------------------------------------------

                chave_preparada = (
                    f"comunicacao_preparada_{id_}"
                )

                comunicacao_preparada_id = (
                    st.session_state.get(
                        chave_preparada
                    )
                )

                if comunicacao_preparada_id:

                    try:

                        comunicacao_preparada = (
                            obter_comunicacao(
                                comunicacao_preparada_id
                            )
                        )

                    except Exception as erro:

                        comunicacao_preparada = None

                        tratar_erro_operacional(
                            erro,
                            "abertura da comunicação preparada",
                        )

                    if (
                        comunicacao_preparada
                        and comunicacao_preparada[
                            "status"
                        ]
                        in [
                            "Rascunho",
                            "Preparada",
                        ]
                    ):

                        st.write("")

                        titulo_secao(
                            t("Outlook Draft")
                        )

                        with st.container(
                            border=True
                        ):

                            st.caption(
                                "DESTINATÁRIOS"
                            )

                            st.text(
                                destinatarios_para_texto(
                                    comunicacao_preparada[
                                        "destinatarios"
                                    ]
                                )
                            )

                            st.caption(
                                "ASSUNTO"
                            )

                            st.write(
                                comunicacao_preparada[
                                    "assunto"
                                ]
                            )

                            st.caption(
                                "MENSAGEM"
                            )

                            st.text(
                                comunicacao_preparada[
                                    "mensagem"
                                ]
                            )

                            eh_encerramento = (
                                comunicacao_preparada.get("tipo") == "Encerramento"
                            )
                            url_outlook = None
                            eml_final = None
                            nome_eml = f"Comunicacao_Encerramento_{codigo}.eml"

                            if eh_encerramento:
                                chave_enc_pdf = f"encerramento_pdf_{id_}"
                                chave_enc_nome = f"encerramento_pdf_nome_{id_}"
                                if chave_enc_pdf in st.session_state:
                                    try:
                                        eml_final = gerar_rascunho_eml_com_anexo(
                                            destinatarios=comunicacao_preparada["destinatarios"],
                                            assunto=comunicacao_preparada["assunto"],
                                            mensagem=comunicacao_preparada["mensagem"],
                                            anexo_bytes=st.session_state[chave_enc_pdf],
                                            nome_anexo=st.session_state.get(
                                                chave_enc_nome, f"Report_Fiscalizacao_Final_{codigo}.pdf"
                                            ),
                                        )
                                    except Exception as erro:
                                        tratar_erro_operacional(
                                            erro, "preparação do e-mail de encerramento"
                                        )
                                else:
                                    st.warning(
                                        "Gere o Report Fiscalizações Final acima para criar o e-mail com anexo automático."
                                    )
                            else:
                                url_outlook = montar_url_outlook_mailto(
                                    destinatarios=comunicacao_preparada["destinatarios"],
                                    assunto=comunicacao_preparada["assunto"],
                                    mensagem=comunicacao_preparada["mensagem"],
                                )

                            oc1, oc2, oc3 = st.columns(3)

                            with oc1:
                                if eh_encerramento and eml_final:
                                    st.download_button(
                                        "Abrir comunicação final (.eml)",
                                        data=eml_final,
                                        file_name=nome_eml,
                                        mime="message/rfc822",
                                        type="primary",
                                        use_container_width=True,
                                    )
                                    st.caption(
                                        "O arquivo contém o Report Fiscalizações Final anexado. Após o download, abra-o no Outlook."
                                    )
                                elif url_outlook:
                                    st.link_button(
                                        t("Abrir no Outlook"),
                                        url_outlook,
                                        type="primary",
                                        use_container_width=True,
                                    )

                            with oc2:

                                if st.button(
                                    t("Marcar como enviada"),
                                    key=(
                                        "confirmar_envio_"
                                        f"{comunicacao_preparada_id}"
                                    ),
                                    use_container_width=True,
                                ):

                                    try:

                                        atualizar_status_comunicacao(
                                            comunicacao_preparada_id,
                                            "Enviada",
                                        )

                                        if comunicacao_preparada.get("tipo") == "Encerramento":
                                            st.session_state.pop(
                                                f"abrir_comunicacao_encerramento_{id_}", None
                                            )

                                        if (
                                            chave_preparada
                                            in st.session_state
                                        ):

                                            del st.session_state[
                                                chave_preparada
                                            ]

                                        exibir_sucesso(
                                            (
                                                "Comunicação marcada "
                                                "como enviada."
                                            )
                                        )

                                        st.rerun()

                                    except Exception as erro:

                                        tratar_erro_operacional(
                                            erro,
                                            "confirmação do envio",
                                        )

                            with oc3:

                                if st.button(
                                    t("Descartar preparação"),
                                    key=(
                                        "descartar_comunicacao_"
                                        f"{comunicacao_preparada_id}"
                                    ),
                                    use_container_width=True,
                                ):

                                    try:

                                        excluir_comunicacao_rascunho(
                                            comunicacao_preparada_id
                                        )

                                        if (
                                            chave_preparada
                                            in st.session_state
                                        ):

                                            del st.session_state[
                                                chave_preparada
                                            ]

                                        exibir_sucesso(
                                            (
                                                "Comunicação preparada "
                                                "foi descartada."
                                            )
                                        )

                                        st.rerun()

                                    except Exception as erro:

                                        tratar_erro_operacional(
                                            erro,
                                            "exclusão da comunicação",
                                        )

                # ---------------------------------------------
                # HISTORICO
                # ---------------------------------------------

                st.write("")

                titulo_secao(
                    t("Communication History")
                )

                try:

                    historico_comunicacoes = (
                        listar_comunicacoes_fiscalizacao(
                            id_
                        )
                    )

                except Exception as erro:

                    historico_comunicacoes = []

                    tratar_erro_operacional(
                        erro,
                        "consulta do histórico de comunicações",
                    )

                if not historico_comunicacoes:

                    st.info(
                        (
                            "Nenhuma comunicação foi registrada "
                            "para esta fiscalização."
                        )
                    )

                else:

                    for resumo_comunicacao in (
                        historico_comunicacoes
                    ):

                        comunicacao_id = (
                            resumo_comunicacao[
                                "id"
                            ]
                        )

                        try:

                            comunicacao_completa = (
                                obter_comunicacao(
                                    comunicacao_id
                                )
                            )

                        except Exception:

                            comunicacao_completa = None

                        with st.container(
                            border=True
                        ):

                            hc1, hc2, hc3 = (
                                st.columns(
                                    [4, 2, 2]
                                )
                            )

                            with hc1:

                                st.caption(
                                    resumo_comunicacao[
                                        "tipo"
                                    ].upper()
                                )

                                st.markdown(
                                    f"**"
                                    f"{resumo_comunicacao['assunto']}"
                                    f"**"
                                )

                                st.caption(
                                    (
                                        "Criada em "
                                        f"{formatar_data_hora(
                                            resumo_comunicacao[
                                                'criado_em'
                                            ]
                                        )}"
                                    )
                                )

                                if (
                                    comunicacao_completa
                                    and comunicacao_completa[
                                        "destinatarios"
                                    ]
                                ):

                                    nomes_emails = [
                                        (
                                            destinatario.get(
                                                "nome"
                                            )
                                            or destinatario.get(
                                                "email"
                                            )
                                            or "-"
                                        )
                                        for destinatario
                                        in comunicacao_completa[
                                            "destinatarios"
                                        ]
                                    ]

                                    st.caption(
                                        (
                                            "Destinatários: "
                                            + ", ".join(
                                                nomes_emails
                                            )
                                        )
                                    )

                            with hc2:

                                st.caption(
                                    "STATUS"
                                )

                                status_comunicacao = (
                                    resumo_comunicacao[
                                        "status"
                                    ]
                                )

                                if (
                                    status_comunicacao
                                    == "Enviada"
                                ):

                                    st.success(
                                        status_comunicacao
                                    )

                                elif (
                                    status_comunicacao
                                    == "Falhou"
                                ):

                                    st.error(
                                        status_comunicacao
                                    )

                                elif (
                                    status_comunicacao
                                    in [
                                        "Preparada",
                                        "Rascunho",
                                    ]
                                ):

                                    st.warning(
                                        status_comunicacao
                                    )

                                else:

                                    st.info(
                                        status_comunicacao
                                    )

                                if resumo_comunicacao[
                                    "enviado_em"
                                ]:

                                    st.caption(
                                        (
                                            "Enviada em "
                                            f"{formatar_data_hora(
                                                resumo_comunicacao[
                                                    'enviado_em'
                                                ]
                                            )}"
                                        )
                                    )

                            with hc3:

                                st.caption(
                                    "DESTINATÁRIOS"
                                )

                                st.metric(
                                    "Total",
                                    resumo_comunicacao[
                                        "quantidade_destinatarios"
                                    ],
                                    label_visibility="collapsed",
                                )

                                if (
                                    status_comunicacao
                                    in [
                                        "Rascunho",
                                        "Preparada",
                                    ]
                                    and comunicacao_completa
                                ):

                                    url_historico = (
                                        montar_url_outlook_mailto(
                                            destinatarios=(
                                                comunicacao_completa[
                                                    "destinatarios"
                                                ]
                                            ),
                                            assunto=(
                                                comunicacao_completa[
                                                    "assunto"
                                                ]
                                            ),
                                            mensagem=(
                                                comunicacao_completa[
                                                    "mensagem"
                                                ]
                                            ),
                                        )
                                    )

                                    if url_historico:

                                        st.link_button(
                                            t("Reabrir no Outlook"),
                                            url_historico,
                                            use_container_width=True,
                                        )


            with aba_flash:
                try:
                    dados_flash = obter_dados_flash_report(id_)
                except Exception as erro:
                    tratar_erro_operacional(erro, "carregamento do Report Fiscalizações")
                    dados_flash = None

                if dados_flash:
                    flash = dados_flash["flash_report"]
                    dados_complementares = dados_flash["dados_complementares"]
                    status_prazo_flash = dados_flash["status_prazo"]
                    documentos_flash = dados_flash.get("documentos") or []
                    tarefa_estracta_flash = (
                        dados_complementares.get("tarefa_estracta") or "-"
                    )

                    st.caption("SYNGENTA BRASIL · TAX CONTROVERSY" if idioma_atual() == "pt" else "SYNGENTA BRAZIL · TAX CONTROVERSY")
                    st.header("Report Fiscalizações")
                    st.caption(f"{codigo} · {titulo}")

                    frs1, frs2, frs3 = st.columns(3)
                    frs1.info(traduzir_valor(flash["status"]))
                    frs2.info(traduzir_valor(fase))
                    if probabilidade == "Provável":
                        frs3.error(traduzir_valor(probabilidade))
                    elif probabilidade == "Possível":
                        frs3.warning(traduzir_valor(probabilidade))
                    else:
                        frs3.success(traduzir_valor(probabilidade))

                    st.divider()
                    titulo_secao(
                        t("Ações do documento"),
                        (
                            ("Gere o Report Fiscalizações consolidado e, quando necessário, o Dossiê Final com todos os documentos originais." if idioma_atual() == "pt" else "Generate the consolidated Tax Inspection Report and, when needed, the Final Case File containing all original documents.")
                        ),
                    )

                    chave_pdf = f"flash_pdf_{id_}"
                    chave_nome = f"flash_pdf_nome_{id_}"
                    chave_zip = f"dossie_zip_{id_}"
                    chave_zip_nome = f"dossie_zip_nome_{id_}"

                    pdf1, pdf2 = st.columns(2)

                    with pdf1:
                        if st.button(
                            t("Gerar Report Fiscalizações consolidado"),
                            key=f"pdf_{id_}",
                            type="primary",
                            use_container_width=True,
                        ):
                            try:
                                pdf_bytes = gerar_flash_report_pdf(
                                    id_,
                                    incorporar_anexos=True,
                                )
                                st.session_state[chave_pdf] = pdf_bytes
                                st.session_state[chave_nome] = (
                                    f"Report_Fiscalizacao_Final_{codigo}.pdf"
                                    if fase == "Encerrada"
                                    else f"Report_Fiscalizacao_{codigo}.pdf"
                                )
                                exibir_sucesso(
                                    "Report Fiscalizações consolidado gerado com sucesso."
                                )
                            except Exception as erro:
                                tratar_erro_operacional(
                                    erro,
                                    "geração do Report Fiscalizações consolidado",
                                )

                        if chave_pdf in st.session_state:
                            st.download_button(
                                t("Baixar Report Fiscalizações"),
                                data=st.session_state[chave_pdf],
                                file_name=st.session_state[chave_nome],
                                mime="application/pdf",
                                use_container_width=True,
                            )

                    with pdf2:
                        if st.button(
                            t("Gerar Dossiê Final ZIP"),
                            key=f"dossie_{id_}",
                            use_container_width=True,
                        ):
                            try:
                                dossie_bytes = gerar_dossie_final_zip(id_)
                                st.session_state[chave_zip] = dossie_bytes
                                st.session_state[chave_zip_nome] = (
                                    f"Dossie_Final_{codigo}.zip"
                                )
                                exibir_sucesso(
                                    "Dossiê Final gerado com sucesso."
                                )
                            except Exception as erro:
                                tratar_erro_operacional(
                                    erro,
                                    "geração do Dossiê Final",
                                )

                        if chave_zip in st.session_state:
                            st.download_button(
                                t("Baixar Dossiê Final"),
                                data=st.session_state[chave_zip],
                                file_name=st.session_state[chave_zip_nome],
                                mime="application/zip",
                                use_container_width=True,
                            )

                    if documentos_flash:
                        st.caption(
                            (
                                f"{len(documentos_flash)} documento(s) no dossiê. "
                                "PDFs e imagens compatíveis são incorporados ao PDF; "
                                "todos os originais permanecem preservados no ZIP."
                            )
                        )

                    st.write("")
                    titulo_secao(t("Report Management"))

                    if esta_arquivada:
                        st.info("Reative o caso para editar o Report Fiscalizações.")
                    else:
                        with st.expander(t("Editar conteúdo do relatório")):
                            with st.form(f"flash_{id_}"):
                                status_flash = st.selectbox(
                                    "Status",
                                    STATUS_FLASH_REPORT,
                                    index=(
                                        STATUS_FLASH_REPORT.index(flash["status"])
                                        if flash["status"] in STATUS_FLASH_REPORT
                                        else 0
                                    ),
                                    format_func=traduzir_opcao,
                                )
                                resumo = st.text_area(
                                    t("Resumo executivo"),
                                    value=flash["resumo_executivo"],
                                    height=150,
                                )
                                problema = st.text_area(
                                    t("Problema / controvérsia"),
                                    value=flash["problema"],
                                    height=120,
                                )
                                causa = st.text_area(
                                    t("Causa / fundamento"),
                                    value=flash["causa"],
                                    height=120,
                                )
                                necessarias = st.text_area(
                                    t("Ações necessárias"),
                                    value=flash["necessarias"],
                                    height=120,
                                )
                                implementadas = st.text_area(
                                    t("Ações implementadas"),
                                    value=flash["implementadas"],
                                    height=120,
                                )
                                riscos = st.text_area(
                                    t("Riscos e impactos"),
                                    value=flash["riscos_impactos"],
                                    height=120,
                                )
                                proximos = st.text_area(
                                    t("Próximos passos"),
                                    value=flash["proximos_passos"],
                                    height=120,
                                )
                                conclusao = st.text_area(
                                    t("Conclusão / posição atual"),
                                    value=flash["conclusao"],
                                    height=120,
                                )
                                salvar_flash = st.form_submit_button(
                                    t("Salvar Report Fiscalizações"),
                                    type="primary",
                                    use_container_width=True,
                                )

                            if salvar_flash:
                                try:
                                    atualizar_flash_report(
                                        fiscalizacao_id=id_,
                                        resumo_executivo=resumo,
                                        problema=problema,
                                        causa=causa,
                                        necessarias=necessarias,
                                        implementadas=implementadas,
                                        riscos_impactos=riscos,
                                        proximos_passos=proximos,
                                        conclusao=conclusao,
                                        status=status_flash,
                                    )
                                    limpar_pdf_sessao(id_)
                                    exibir_sucesso("Report Fiscalizações atualizado.")
                                    st.rerun()
                                except Exception as erro:
                                    tratar_erro_operacional(
                                        erro,
                                        "atualização do Report Fiscalizações",
                                    )

                    st.write("")
                    titulo_secao(t("01 — Identificação"))
                    ri1, ri2, ri3, ri4 = st.columns(4)
                    with ri1:
                        with st.container(border=True):
                            st.caption("EMPRESA")
                            st.write(empresa)
                            st.caption("FILIAL")
                            st.write(filial or "-")
                    with ri2:
                        with st.container(border=True):
                            st.caption("ÓRGÃO")
                            st.write(orgao or "-")
                            st.caption("RESPONSÁVEL")
                            st.write(responsavel_principal or "-")
                    with ri3:
                        with st.container(border=True):
                            st.caption("FASE")
                            st.write(traduzir_valor(fase))
                            st.caption("PROBABILIDADE")
                            st.write(traduzir_valor(probabilidade))
                    with ri4:
                        with st.container(border=True):
                            st.caption("EXPOSIÇÃO")
                            st.write(formatar_moeda(exposicao))
                            st.caption("TRIBUTOS")
                            st.write(", ".join(tributos) if tributos else "-")

                    st.write("")
                    ref1, ref2 = st.columns(2)
                    with ref1:
                        with st.container(border=True):
                            st.caption("TAREFA E-STRACTA")
                            st.write(tarefa_estracta_flash)
                    with ref2:
                        with st.container(border=True):
                            st.caption("DOCUMENTOS NO DOSSIÊ")
                            st.write(len(documentos_flash))

                    titulo_secao(t("02 — Executive Summary"))
                    with st.container(border=True):
                        exibir_texto_relatorio(flash["resumo_executivo"])

                    titulo_secao(t("03 — Inspection Background"))
                    with st.container(border=True):
                        st.markdown("**Objeto / Escopo**")
                        exibir_texto_relatorio(objeto)

                    titulo_secao(t("04 — Tax Controversy"))
                    rc1, rc2 = st.columns(2)
                    with rc1:
                        with st.container(border=True):
                            st.markdown("**Problema / Controvérsia**")
                            exibir_texto_relatorio(flash["problema"])
                    with rc2:
                        with st.container(border=True):
                            st.markdown("**Causa / Fundamento**")
                            exibir_texto_relatorio(flash["causa"])

                    titulo_secao(t("05 — Developments"))
                    if not historico:
                        st.info(t(t("Nenhum desdobramento registrado.")))
                    else:
                        for item in reversed(historico):
                            hist_id, hist_fase, hist_data, hist_obs, hist_criado = item
                            with st.container(border=True):
                                st.caption(formatar_data(hist_data))
                                st.markdown(f"**{traduzir_valor(hist_fase)}**")
                                if hist_obs:
                                    st.write(hist_obs)

                    titulo_secao(t("06 — Deadlines"))
                    dl1, dl2, dl3 = st.columns(3)
                    dl1.metric(
                        t("Recebimento"),
                        formatar_data(dados_complementares["data_recebimento"]),
                    )
                    dl2.metric(
                        t("Prazo atual"),
                        formatar_data(dados_complementares["prazo_resposta"]),
                    )
                    with dl3:
                        exibir_status_prazo(
                            status_prazo_flash["status"],
                            status_prazo_flash["dias_restantes"],
                        )

                    titulo_secao(t("07 — Actions Taken"))
                    aa1, aa2 = st.columns(2)
                    with aa1:
                        with st.container(border=True):
                            st.markdown("**Ações necessárias**")
                            exibir_texto_relatorio(flash["necessarias"])
                    with aa2:
                        with st.container(border=True):
                            st.markdown("**Ações implementadas**")
                            exibir_texto_relatorio(flash["implementadas"])

                    titulo_secao(t("08 — Internal Tasks"))
                    if not tarefas:
                        st.info(t(t("Nenhuma tarefa cadastrada.")))
                    else:
                        dados_tarefas = [
                            {
                                "Atividade": tarefa[1],
                                "Responsável": tarefa[3] or "-",
                                "Prazo": formatar_data(tarefa[4]),
                                "Status" if idioma_atual() == "pt" else "Status": traduzir_valor(tarefa[5]),
                            }
                            for tarefa in tarefas
                        ]
                        st.dataframe(
                            dados_tarefas,
                            use_container_width=True,
                            hide_index=True,
                        )

                    titulo_secao(t("09 — Documents & Evidence"))
                    if not documentos_flash:
                        st.info(
                            "Nenhum documento ou evidência foi registrado no dossiê."
                        )
                    else:
                        st.caption(
                            f"{len(documentos_flash)} documento(s) registrado(s) no dossiê eletrônico."
                        )
                        for indice_documento, documento in enumerate(
                            documentos_flash,
                            start=1,
                        ):
                            with st.container(border=True):
                                ed1, ed2 = st.columns([4, 1.5])
                                with ed1:
                                    st.caption(documento["tipo_documento"].upper())
                                    st.markdown(
                                        f"**{indice_documento:02d} · {documento['nome_original']}**"
                                    )
                                    if documento.get("descricao"):
                                        st.write(documento["descricao"])
                                    dados_documento = []
                                    if documento.get("data_documento"):
                                        dados_documento.append(
                                            "Data do documento: "
                                            + formatar_data(documento["data_documento"])
                                        )
                                    dados_documento.append(
                                        formatar_tamanho_arquivo(
                                            documento.get("tamanho_bytes")
                                        )
                                    )
                                    st.caption(" · ".join(dados_documento))
                                with ed2:
                                    st.caption("INTEGRIDADE")
                                    try:
                                        integridade_flash = verificar_integridade_documento(
                                            documento["id"]
                                        )
                                        if integridade_flash["ok"]:
                                            st.success("Arquivo íntegro")
                                        else:
                                            st.error(integridade_flash["motivo"])
                                    except Exception:
                                        st.warning("Não verificada")
                                    hash_documento = documento.get("hash_sha256") or ""
                                    if hash_documento:
                                        st.caption("SHA-256")
                                        st.code(hash_documento[:16] + "...", language=None)

                    titulo_secao(t("10 — Risks & Impacts"))
                    with st.container(border=True):
                        exibir_texto_relatorio(flash["riscos_impactos"])

                    titulo_secao(t("11 — Next Steps"))
                    with st.container(border=True):
                        exibir_texto_relatorio(flash["proximos_passos"])

                    titulo_secao(t("12 — Current Position"))
                    with st.container(border=True):
                        exibir_texto_relatorio(flash["conclusao"])

                    st.divider()
                    rod1, rod2, rod3 = st.columns(3)
                    with rod1:
                        st.caption("ÚLTIMA ATUALIZAÇÃO")
                        st.write(formatar_data_hora(flash["atualizado_em"]))
                    with rod2:
                        st.caption("TAREFA E-STRACTA")
                        st.write(tarefa_estracta_flash)
                    with rod3:
                        st.caption("DOCUMENTOS NO DOSSIÊ")
                        st.write(len(documentos_flash))


# ============================================================
# AREAS & ENVOLVIDOS
# FASE 2.0A2
# ============================================================

elif st.session_state.pagina == "Áreas & Envolvidos":
    titulo_pagina(
        t("Áreas & Envolvidos"),
        (
            "Cadastro das áreas e pessoas que participam das fiscalizações e que poderão receber comunicações do Fiscal Tracker."
            if idioma_atual() == "pt"
            else "Register the areas and people involved in tax inspections and eligible to receive Fiscal Tracker communications."
        ),
    )

    try:
        areas_admin = listar_areas_cadastradas(
            incluir_inativas=True
        )
        pessoas_admin = listar_pessoas_cadastradas(
            incluir_inativas=True
        )
    except Exception as erro:
        tratar_erro_operacional(
            erro,
            "consulta de áreas e envolvidos",
        )
        areas_admin = []
        pessoas_admin = []

    areas_ativas_admin = [
        area
        for area in areas_admin
        if area["ativa"]
    ]

    pessoas_ativas_admin = [
        pessoa
        for pessoa in pessoas_admin
        if pessoa["ativa"]
    ]

    a1, a2, a3 = st.columns(3)

    with a1:
        st.metric(
            t("Áreas ativas"),
            len(
                areas_ativas_admin
            ),
        )

    with a2:
        st.metric(
            t("Pessoas ativas"),
            len(
                pessoas_ativas_admin
            ),
        )

    with a3:
        total_vinculos = 0

        for area in areas_admin:
            try:
                total_vinculos += len(
                    listar_pessoas_area(
                        area["id"]
                    )
                )
            except Exception:
                pass

        st.metric(
            t("Vínculos"),
            total_vinculos,
        )

    st.write("")

    aba_areas, aba_pessoas, aba_vinculos = st.tabs(
        [
            t("Áreas"),
            t("Pessoas"),
            t("Vínculos"),
        ]
    )

    # ========================================================
    # ABA ÁREAS
    # ========================================================

    with aba_areas:
        titulo_secao(
            t("Cadastro de Áreas"),
            (
                "Mantenha aqui as áreas que podem participar "
                "das fiscalizações."
            ),
        )

        with st.expander(
            t("Cadastrar nova área"),
            expanded=(
                len(
                    areas_admin
                )
                == 0
            ),
        ):
            with st.form(
                "form_nova_area",
                clear_on_submit=True,
            ):
                nova_area_nome = st.text_input(
                    "Nome da área *",
                    placeholder=(
                        "Ex.: Jurídico Tributário"
                    ),
                )

                nova_area_descricao = st.text_area(
                    "Descrição",
                    placeholder=(
                        "Breve descrição da atuação da área."
                    ),
                    height=90,
                )

                nova_area_email = st.text_input(
                    "E-mail da área",
                    placeholder=(
                        "Ex.: tax.legal@empresa.com"
                    ),
                    help=(
                        "Opcional. Poderá ser utilizado "
                        "nas comunicações futuras."
                    ),
                )

                salvar_nova_area = (
                    st.form_submit_button(
                        t("Cadastrar área"),
                        type="primary",
                        use_container_width=True,
                    )
                )

            if salvar_nova_area:
                erros_area = (
                    validar_campos_obrigatorios(
                        [
                            (
                                "o nome da área",
                                nova_area_nome,
                            ),
                        ]
                    )
                )

                if erros_area:
                    for erro in erros_area:
                        st.warning(
                            erro
                        )
                else:
                    try:
                        cadastrar_area(
                            nome=(
                                nova_area_nome.strip()
                            ),
                            descricao=(
                                nova_area_descricao
                            ),
                            email_area=(
                                nova_area_email
                            ),
                        )

                        exibir_sucesso(
                            "Área cadastrada com sucesso."
                        )

                        st.rerun()

                    except Exception as erro:
                        tratar_erro_operacional(
                            erro,
                            "criação da área",
                        )

        st.write("")
        titulo_secao(
            t("Áreas cadastradas")
        )

        if not areas_admin:
            st.info(
                "Nenhuma área cadastrada."
            )

        else:
            for area in areas_admin:
                area_id = area["id"]

                pessoas_area = []

                try:
                    pessoas_area = (
                        listar_pessoas_area(
                            area_id
                        )
                    )
                except Exception:
                    pass

                status_texto = (
                    "ATIVA"
                    if area["ativa"]
                    else "INATIVA"
                )

                with st.expander(
                    f"{area['nome']} · {status_texto}"
                ):
                    ac1, ac2, ac3 = (
                        st.columns(
                            [3, 3, 2]
                        )
                    )

                    with ac1:
                        st.caption(
                            "E-MAIL DA ÁREA"
                        )

                        st.write(
                            area[
                                "email_area"
                            ]
                            or "-"
                        )

                    with ac2:
                        st.caption(
                            "ENVOLVIDOS VINCULADOS"
                        )

                        st.write(
                            len(
                                pessoas_area
                            )
                        )

                    with ac3:
                        st.caption(
                            "STATUS"
                        )

                        if area["ativa"]:
                            st.success(
                                "Ativa"
                            )
                        else:
                            st.warning(
                                "Inativa"
                            )

                    if area[
                        "descricao"
                    ]:
                        st.caption(
                            "DESCRIÇÃO"
                        )
                        st.write(
                            area[
                                "descricao"
                            ]
                        )

                    st.divider()

                    with st.form(
                        f"editar_area_{area_id}"
                    ):
                        editar_area_nome = (
                            st.text_input(
                                "Nome",
                                value=(
                                    area[
                                        "nome"
                                    ]
                                ),
                                key=(
                                    f"nome_area_"
                                    f"{area_id}"
                                ),
                            )
                        )

                        editar_area_descricao = (
                            st.text_area(
                                "Descrição",
                                value=(
                                    area[
                                        "descricao"
                                    ]
                                    or ""
                                ),
                                height=90,
                                key=(
                                    f"desc_area_"
                                    f"{area_id}"
                                ),
                            )
                        )

                        editar_area_email = (
                            st.text_input(
                                "E-mail da área",
                                value=(
                                    area[
                                        "email_area"
                                    ]
                                    or ""
                                ),
                                key=(
                                    f"email_area_"
                                    f"{area_id}"
                                ),
                            )
                        )

                        salvar_edicao_area = (
                            st.form_submit_button(
                                t("Salvar alterações"),
                                use_container_width=True,
                            )
                        )

                    if salvar_edicao_area:
                        try:
                            atualizar_area(
                                area_id=area_id,
                                nome=(
                                    editar_area_nome
                                ),
                                descricao=(
                                    editar_area_descricao
                                ),
                                email_area=(
                                    editar_area_email
                                ),
                            )

                            exibir_sucesso(
                                "Área atualizada."
                            )

                            st.rerun()

                        except Exception as erro:
                            tratar_erro_operacional(
                                erro,
                                "atualização da área",
                            )

                    if area["ativa"]:
                        if st.button(
                            "Desativar área",
                            key=(
                                f"desativar_area_"
                                f"{area_id}"
                            ),
                            use_container_width=True,
                        ):
                            try:
                                alterar_status_area(
                                    area_id,
                                    False,
                                )

                                exibir_sucesso(
                                    "Área desativada."
                                )

                                st.rerun()

                            except Exception as erro:
                                tratar_erro_operacional(
                                    erro,
                                    "desativação da área",
                                )
                    else:
                        if st.button(
                            "Reativar área",
                            key=(
                                f"reativar_area_"
                                f"{area_id}"
                            ),
                            type="primary",
                            use_container_width=True,
                        ):
                            try:
                                alterar_status_area(
                                    area_id,
                                    True,
                                )

                                exibir_sucesso(
                                    "Área reativada."
                                )

                                st.rerun()

                            except Exception as erro:
                                tratar_erro_operacional(
                                    erro,
                                    "reativação da área",
                                )

    # ========================================================
    # ABA PESSOAS
    # ========================================================

    with aba_pessoas:
        titulo_secao(
            t("Cadastro de Pessoas"),
            (
                "Cadastre os profissionais que poderão ser "
                "associados às áreas e às comunicações."
            ),
        )

        with st.expander(
            t("Cadastrar nova pessoa"),
            expanded=(
                len(
                    pessoas_admin
                )
                == 0
            ),
        ):
            with st.form(
                "form_nova_pessoa",
                clear_on_submit=True,
            ):
                nova_pessoa_nome = (
                    st.text_input(
                        "Nome *",
                        placeholder=(
                            "Ex.: Maria Silva"
                        ),
                    )
                )

                nova_pessoa_email = (
                    st.text_input(
                        "E-mail *",
                        placeholder=(
                            "Ex.: maria.silva@empresa.com"
                        ),
                    )
                )

                nova_pessoa_cargo = (
                    st.text_input(
                        "Cargo / função",
                        placeholder=(
                            "Ex.: Tax Manager"
                        ),
                    )
                )

                salvar_nova_pessoa = (
                    st.form_submit_button(
                        t("Cadastrar pessoa"),
                        type="primary",
                        use_container_width=True,
                    )
                )

            if salvar_nova_pessoa:
                erros_pessoa = (
                    validar_campos_obrigatorios(
                        [
                            (
                                "o nome da pessoa",
                                nova_pessoa_nome,
                            ),
                            (
                                "o e-mail da pessoa",
                                nova_pessoa_email,
                            ),
                        ]
                    )
                )

                if erros_pessoa:
                    for erro in erros_pessoa:
                        st.warning(
                            erro
                        )
                else:
                    try:
                        cadastrar_pessoa(
                            nome=(
                                nova_pessoa_nome
                            ),
                            email=(
                                nova_pessoa_email
                            ),
                            cargo=(
                                nova_pessoa_cargo
                            ),
                        )

                        exibir_sucesso(
                            "Pessoa cadastrada com sucesso."
                        )

                        st.rerun()

                    except Exception as erro:
                        tratar_erro_operacional(
                            erro,
                            "criação da pessoa",
                        )

        st.write("")
        titulo_secao(
            t("Pessoas cadastradas")
        )

        if not pessoas_admin:
            st.info(
                "Nenhuma pessoa cadastrada."
            )

        else:
            for pessoa in pessoas_admin:
                pessoa_id = pessoa["id"]

                status_texto = (
                    "ATIVA"
                    if pessoa["ativa"]
                    else "INATIVA"
                )

                with st.expander(
                    f"{pessoa['nome']} · {status_texto}"
                ):
                    pc1, pc2, pc3 = (
                        st.columns(
                            [3, 3, 2]
                        )
                    )

                    with pc1:
                        st.caption(
                            "E-MAIL"
                        )
                        st.write(
                            pessoa[
                                "email"
                            ]
                        )

                    with pc2:
                        st.caption(
                            "CARGO / FUNÇÃO"
                        )
                        st.write(
                            pessoa[
                                "cargo"
                            ]
                            or "-"
                        )

                    with pc3:
                        st.caption(
                            "STATUS"
                        )

                        if pessoa["ativa"]:
                            st.success(
                                "Ativa"
                            )
                        else:
                            st.warning(
                                "Inativa"
                            )

                    try:
                        areas_da_pessoa = []

                        for area_item in (
                            areas_admin
                        ):
                            pessoas_area = (
                                listar_pessoas_area(
                                    area_item[
                                        "id"
                                    ]
                                )
                            )

                            if any(
                                item[
                                    "id"
                                ]
                                == pessoa_id
                                for item
                                in pessoas_area
                            ):
                                areas_da_pessoa.append(
                                    area_item[
                                        "nome"
                                    ]
                                )

                        st.caption(
                            "ÁREAS"
                        )

                        st.write(
                            ", ".join(
                                areas_da_pessoa
                            )
                            if areas_da_pessoa
                            else "-"
                        )

                    except Exception:
                        pass

                    st.divider()

                    with st.form(
                        f"editar_pessoa_{pessoa_id}"
                    ):
                        editar_pessoa_nome = (
                            st.text_input(
                                "Nome",
                                value=(
                                    pessoa[
                                        "nome"
                                    ]
                                ),
                                key=(
                                    f"nome_pessoa_"
                                    f"{pessoa_id}"
                                ),
                            )
                        )

                        editar_pessoa_email = (
                            st.text_input(
                                "E-mail",
                                value=(
                                    pessoa[
                                        "email"
                                    ]
                                ),
                                key=(
                                    f"email_pessoa_"
                                    f"{pessoa_id}"
                                ),
                            )
                        )

                        editar_pessoa_cargo = (
                            st.text_input(
                                "Cargo / função",
                                value=(
                                    pessoa[
                                        "cargo"
                                    ]
                                    or ""
                                ),
                                key=(
                                    f"cargo_pessoa_"
                                    f"{pessoa_id}"
                                ),
                            )
                        )

                        salvar_edicao_pessoa = (
                            st.form_submit_button(
                                t("Salvar alterações"),
                                use_container_width=True,
                            )
                        )

                    if salvar_edicao_pessoa:
                        try:
                            atualizar_pessoa(
                                pessoa_id=(
                                    pessoa_id
                                ),
                                nome=(
                                    editar_pessoa_nome
                                ),
                                email=(
                                    editar_pessoa_email
                                ),
                                cargo=(
                                    editar_pessoa_cargo
                                ),
                            )

                            exibir_sucesso(
                                "Pessoa atualizada."
                            )

                            st.rerun()

                        except Exception as erro:
                            tratar_erro_operacional(
                                erro,
                                "atualização da pessoa",
                            )

                    if pessoa["ativa"]:
                        if st.button(
                            "Desativar pessoa",
                            key=(
                                f"desativar_pessoa_"
                                f"{pessoa_id}"
                            ),
                            use_container_width=True,
                        ):
                            try:
                                alterar_status_pessoa(
                                    pessoa_id,
                                    False,
                                )

                                exibir_sucesso(
                                    "Pessoa desativada."
                                )

                                st.rerun()

                            except Exception as erro:
                                tratar_erro_operacional(
                                    erro,
                                    "desativação da pessoa",
                                )
                    else:
                        if st.button(
                            "Reativar pessoa",
                            key=(
                                f"reativar_pessoa_"
                                f"{pessoa_id}"
                            ),
                            type="primary",
                            use_container_width=True,
                        ):
                            try:
                                alterar_status_pessoa(
                                    pessoa_id,
                                    True,
                                )

                                exibir_sucesso(
                                    "Pessoa reativada."
                                )

                                st.rerun()

                            except Exception as erro:
                                tratar_erro_operacional(
                                    erro,
                                    "reativação da pessoa",
                                )

    # ========================================================
    # ABA VÍNCULOS
    # ========================================================

    with aba_vinculos:
        titulo_secao(
            t("Vínculos Área × Pessoa"),
            (
                "Defina quem participa de cada área e o papel "
                "de cada envolvido."
            ),
        )

        if (
            not areas_ativas_admin
            or not pessoas_ativas_admin
        ):
            st.info(
                "Para criar vínculos, mantenha pelo menos "
                "uma área ativa e uma pessoa ativa."
            )

        else:
            mapa_areas = {
                area["nome"]: area
                for area in areas_ativas_admin
            }

            mapa_pessoas = {
                (
                    f"{pessoa['nome']} "
                    f"· {pessoa['email']}"
                ): pessoa
                for pessoa
                in pessoas_ativas_admin
            }

            with st.form(
                "form_vinculo_area_pessoa"
            ):
                vinculo_area_nome = (
                    st.selectbox(
                        "Área *",
                        list(
                            mapa_areas.keys()
                        ),
                    )
                )

                vinculo_pessoa_label = (
                    st.selectbox(
                        "Pessoa *",
                        list(
                            mapa_pessoas.keys()
                        ),
                    )
                )

                vinculo_papel = (
                    st.selectbox(
                        "Papel",
                        PAPEIS_ENVOLVIMENTO,
                        index=(
                            PAPEIS_ENVOLVIMENTO.index(
                                "Informado"
                            )
                            if "Informado"
                            in PAPEIS_ENVOLVIMENTO
                            else 0
                        ),
                    )
                )

                vinculo_principal = (
                    st.checkbox(
                        "Responsável principal da área",
                        help=(
                            "Ao marcar esta opção, qualquer "
                            "responsável principal anterior "
                            "da mesma área será desmarcado."
                        ),
                    )
                )

                salvar_vinculo = (
                    st.form_submit_button(
                        t("Salvar vínculo"),
                        type="primary",
                        use_container_width=True,
                    )
                )

            if salvar_vinculo:
                try:
                    vincular_pessoa_area(
                        area_id=(
                            mapa_areas[
                                vinculo_area_nome
                            ][
                                "id"
                            ]
                        ),
                        pessoa_id=(
                            mapa_pessoas[
                                vinculo_pessoa_label
                            ][
                                "id"
                            ]
                        ),
                        papel=(
                            vinculo_papel
                        ),
                        principal=(
                            vinculo_principal
                        ),
                    )

                    exibir_sucesso(
                        "Vínculo salvo com sucesso."
                    )

                    st.rerun()

                except Exception as erro:
                    tratar_erro_operacional(
                        erro,
                        "criação do vínculo",
                    )

        st.write("")
        titulo_secao(
            t("Mapa de Envolvidos")
        )

        if not areas_admin:
            st.info(
                "Nenhuma área cadastrada."
            )

        else:
            for area in areas_admin:
                try:
                    pessoas_area = (
                        listar_pessoas_area(
                            area[
                                "id"
                            ]
                        )
                    )
                except Exception as erro:
                    tratar_erro_operacional(
                        erro,
                        "consulta dos vínculos",
                    )
                    pessoas_area = []

                with st.container(
                    border=True
                ):
                    mc1, mc2 = (
                        st.columns(
                            [4, 2]
                        )
                    )

                    with mc1:
                        st.markdown(
                            f"### {area['nome']}"
                        )

                        if area[
                            "descricao"
                        ]:
                            st.caption(
                                area[
                                    "descricao"
                                ]
                            )

                        if area[
                            "email_area"
                        ]:
                            st.write(
                                f"**E-mail da área:** "
                                f"{area['email_area']}"
                            )

                    with mc2:
                        if area["ativa"]:
                            st.success(
                                "Área ativa"
                            )
                        else:
                            st.warning(
                                "Área inativa"
                            )

                    if not pessoas_area:
                        st.caption(
                            "Nenhuma pessoa vinculada."
                        )

                    else:
                        for pessoa_vinculada in pessoas_area:
                            vc1, vc2, vc3, vc4 = (
                                st.columns(
                                    [4, 2, 2, 1]
                                )
                            )

                            with vc1:
                                destaque = (
                                    " · PRINCIPAL"
                                    if pessoa_vinculada[
                                        "principal"
                                    ]
                                    else ""
                                )

                                st.write(
                                    f"**"
                                    f"{pessoa_vinculada['nome']}"
                                    f"{destaque}"
                                    f"**"
                                )

                                st.caption(
                                    pessoa_vinculada[
                                        "email"
                                    ]
                                )

                            with vc2:
                                st.caption(
                                    "PAPEL"
                                )
                                st.write(
                                    pessoa_vinculada[
                                        "papel"
                                    ]
                                )

                            with vc3:
                                st.caption(
                                    "CARGO / FUNÇÃO"
                                )
                                st.write(
                                    pessoa_vinculada[
                                        "cargo"
                                    ]
                                    or "-"
                                )

                            with vc4:
                                if st.button(
                                    "Remover",
                                    key=(
                                        f"remover_vinculo_"
                                        f"{area['id']}_"
                                        f"{pessoa_vinculada['id']}"
                                    ),
                                    use_container_width=True,
                                ):
                                    try:
                                        remover_vinculo_pessoa_area(
                                            area[
                                                "id"
                                            ],
                                            pessoa_vinculada[
                                                "id"
                                            ],
                                        )

                                        exibir_sucesso(
                                            "Vínculo removido."
                                        )

                                        st.rerun()

                                    except Exception as erro:
                                        tratar_erro_operacional(
                                            erro,
                                            "remoção do vínculo",
                                        )

                    try:
                        destinatarios = (
                            obter_destinatarios_area(
                                area[
                                    "id"
                                ]
                            )
                        )

                        st.divider()
                        st.caption(
                            "PRÉVIA DE DESTINATÁRIOS "
                            "PARA A FASE 2.0B"
                        )

                        if destinatarios:
                            st.write(
                                ", ".join(
                                    item[
                                        "email"
                                    ]
                                    for item
                                    in destinatarios
                                )
                            )
                        else:
                            st.write(
                                "Nenhum destinatário ativo."
                            )

                    except Exception:
                        pass


# ============================================================
# NOVA FISCALIZACAO
# ============================================================

elif st.session_state.pagina == "Nova fiscalização":
    titulo_pagina(
        t("Nova fiscalização"),
        ("Registro de novo caso no portfólio Tax Controversy." if idioma_atual() == "pt" else "Register a new case in the Tax Controversy portfolio."),
    )

    with st.form("nova_fiscalizacao", clear_on_submit=True):
        st.subheader(t("Case Information"))
        titulo = st.text_input("Título *")
        objeto = st.text_area("Objeto / escopo", height=120)
        nf1, nf2 = st.columns(2)
        with nf1:
            empresa = st.text_input("Empresa *")
        with nf2:
            filial = st.text_input("Filial")
        responsavel = st.text_input("Responsável principal")
        tarefa_estracta = st.text_input(
            "Tarefa e-Stracta",
            help="Número ou código da tarefa correspondente na ferramenta interna.",
        )

        st.subheader(t("Deadlines"))
        nd1, nd2 = st.columns(2)
        with nd1:
            recebimento = st.date_input(
                "Data de recebimento",
                value=date.today(),
                format="DD/MM/YYYY",
            )
        with nd2:
            prazo = st.date_input(
                "Prazo de resposta",
                value=date.today(),
                format="DD/MM/YYYY",
            )

        st.subheader(t("Tax Information"))
        ni1, ni2 = st.columns(2)
        with ni1:
            orgao = st.text_input("Órgão fiscalizador")
        with ni2:
            probabilidade = st.selectbox(
                t("Probabilidade"),
                PROBABILIDADES,
                index=1,
                format_func=traduzir_opcao,
            )

        exposicao = st.number_input(
            "Exposição estimada (R$)",
            min_value=0.0,
            value=0.0,
            step=1000.0,
        )
        tributos = st.multiselect("Tributos envolvidos", TRIBUTOS_DISPONIVEIS)
        areas = st.multiselect(
            "Áreas envolvidas",
            obter_nomes_areas_ativas(),
        )

        st.subheader(
            "Flash Report Executivo"
            if idioma_atual() == "pt"
            else "Executive Flash Report"
        )
        st.caption(
            (
                "Use esta classificação apenas para fiscalizações com relevância executiva. "
                "Ela controlará a inclusão no Flash Report do Dashboard."
            )
            if idioma_atual() == "pt"
            else (
                "Use this classification only for executive-relevant tax inspections. "
                "It controls inclusion in the Dashboard Flash Report."
            )
        )
        incluir_flash_report = st.checkbox(
            (
                "Incluir esta fiscalização no Flash Report"
                if idioma_atual() == "pt"
                else "Include this tax inspection in the Flash Report"
            ),
            value=False,
        )
        prioridade_executiva = st.selectbox(
            (
                "Prioridade executiva"
                if idioma_atual() == "pt"
                else "Executive priority"
            ),
            PRIORIDADES_EXECUTIVAS,
            index=1,
            format_func=traduzir_opcao,
            help=(
                "Alta, Média ou Baixa. Será usada na ordenação do Flash Report."
                if idioma_atual() == "pt"
                else "High, Medium or Low. Used to order the Flash Report."
            ),
        )
        atualizacao_executiva = st.text_area(
            (
                "Atualização executiva"
                if idioma_atual() == "pt"
                else "Executive update"
            ),
            height=100,
            max_chars=500,
            help=(
                "Texto curto para o Flash Report. Pode ser preenchido depois. "
                "Se vazio, o Resumo Executivo do Report Fiscalizações poderá ser usado como fallback."
                if idioma_atual() == "pt"
                else (
                    "Short text for the Flash Report. It can be completed later. "
                    "If blank, the Tax Inspection Report Executive Summary may be used as fallback."
                )
            ),
        )

        salvar = st.form_submit_button(
            t("Cadastrar fiscalização"),
            type="primary",
            use_container_width=True,
        )

    if salvar:
        erros = validar_campos_obrigatorios(
            [
                ("o título da fiscalização", titulo),
                ("a empresa", empresa),
            ]
        )
        erro_periodo = validar_periodo(
            recebimento,
            prazo,
            "data de recebimento",
            "data de resposta",
        )
        if erro_periodo:
            erros.append(erro_periodo)

        if erros:
            for erro in erros:
                st.warning(erro)
        else:
            try:
                codigo = cadastrar_fiscalizacao(
                    titulo=titulo.strip(),
                    objeto=objeto.strip(),
                    empresa=empresa.strip(),
                    filial=filial.strip(),
                    orgao=orgao.strip(),
                    exposicao=exposicao,
                    probabilidade=probabilidade,
                    tributos=tributos,
                    areas=areas,
                    data_recebimento=recebimento,
                    prazo_resposta=prazo,
                    responsavel_principal=responsavel.strip(),
                    tarefa_estracta=tarefa_estracta.strip(),
                    incluir_flash_report=incluir_flash_report,
                    prioridade_executiva=prioridade_executiva,
                    atualizacao_executiva=atualizacao_executiva.strip(),
                )
                exibir_sucesso(
                    f"Fiscalização {codigo} cadastrada com sucesso."
                )
            except Exception as erro:
                tratar_erro_operacional(erro, "criação da fiscalização")
